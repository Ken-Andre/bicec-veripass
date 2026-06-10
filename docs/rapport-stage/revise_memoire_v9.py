from __future__ import annotations

import re
import shutil
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Inches
from docx.text.paragraph import Paragraph
from PIL import Image, ImageDraw


ROOT = Path(__file__).resolve().parents[2]
SRC = ROOT / "docs" / "rapport-stage" / "memoire-stage-andre-yoann-kenmogne-v8.docx"
OUT = ROOT / "docs" / "rapport-stage" / "memoire-stage-andre-yoann-kenmogne-v9.docx"
FIG_DIR = ROOT / "docs" / "rapport-stage" / "figures"
MOBILE_SCOPE = ROOT / "docs" / "test-evidence" / "latest" / "kyc-compliance-demo" / "screens" / "11-mobile-document-scope-cni-only-disabled-options.png"
BACKOFFICE_DASHBOARD = ROOT / "docs" / "test-evidence" / "latest" / "backoffice" / "screens" / "dashboard.png"
BACKOFFICE_REDACTED = FIG_DIR / "backoffice_dashboard_redacted_v9.png"


def set_text(paragraph: Paragraph, text: str) -> None:
    style = paragraph.style
    for run in paragraph.runs:
        run.text = ""
    paragraph.add_run(text)
    paragraph.style = style


def insert_after(paragraph: Paragraph, text: str = "", style: str | None = None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    p = Paragraph(new_p, paragraph._parent)
    if style:
        p.style = style
    if text:
        p.add_run(text)
    return p


def insert_paragraphs_after(paragraph: Paragraph, paragraphs: list[str], style: str = "Memoire Body") -> Paragraph:
    cur = paragraph
    for text in paragraphs:
        cur = insert_after(cur, text, style)
    return cur


def find_contains(doc: Document, needle: str) -> Paragraph:
    for p in doc.paragraphs:
        if needle in p.text:
            return p
    raise ValueError(f"Paragraph containing {needle!r} not found")


def find_exact(doc: Document, text: str) -> Paragraph:
    for p in doc.paragraphs:
        if p.text.strip() == text:
            return p
    raise ValueError(f"Paragraph {text!r} not found")


def replace_exact(doc: Document, old: str, new: str) -> None:
    p = find_exact(doc, old)
    set_text(p, new)


def replace_contains(doc: Document, needle: str, new: str) -> None:
    p = find_contains(doc, needle)
    set_text(p, new)


def style_name(doc: Document, prefix: str, fallback: str = "Memoire Body") -> str:
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefix):
            return p.style.name
    return fallback


def add_abbreviation_rows(doc: Document) -> None:
    table = None
    for t in doc.tables:
        if not t.rows:
            continue
        header = [c.text.strip() for c in t.rows[0].cells]
        if header[:2] == ["Abréviation", "Signification"]:
            table = t
            break
    if table is None:
        raise ValueError("Abbreviation table not found")

    existing = {row.cells[0].text.strip() for row in table.rows[1:]}
    additions = [
        ("CGAP", "Consultative Group to Assist the Poor, organisme de référence sur la finance inclusive"),
        ("GAFI", "Groupe d'action financière, organisme international de lutte contre le blanchiment"),
        ("JSONB", "Format JSON binaire de PostgreSQL, utile pour les métadonnées techniques variables"),
        ("OLAP", "Online Analytical Processing, traitement analytique orienté agrégats et décisionnel"),
        ("OLTP", "Online Transaction Processing, traitement transactionnel orienté opérations métier"),
        ("TTL", "Time To Live, durée de validité d'une donnée volatile en cache"),
    ]
    for abbr, meaning in additions:
        if abbr in existing:
            continue
        row = table.add_row()
        row.cells[0].text = abbr
        row.cells[1].text = meaning


def make_backoffice_redacted() -> None:
    if not BACKOFFICE_DASHBOARD.exists():
        return
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    image = Image.open(BACKOFFICE_DASHBOARD).convert("RGB")
    w, h = image.size
    image = image.crop((int(w * 0.18), int(h * 0.06), w, int(h * 0.82)))
    draw = ImageDraw.Draw(image)
    cw, ch = image.size
    draw.rectangle((int(cw * 0.70), 0, cw, int(ch * 0.10)), fill=(255, 255, 255))
    image.save(BACKOFFICE_REDACTED)


def insert_figure_after(
    paragraph: Paragraph,
    image_path: Path,
    caption: str,
    source: str,
    width_inches: float,
) -> Paragraph:
    if not image_path.exists():
        return paragraph
    image_p = insert_after(paragraph, "", "Memoire Body")
    image_p.alignment = 1
    image_p.add_run().add_picture(str(image_path), width=Inches(width_inches))
    caption_p = insert_after(image_p, caption, "Caption Memoire")
    source_p = insert_after(caption_p, source, "Source Memoire")
    return source_p


def update_front_lists(doc: Document) -> None:
    replace_exact(doc, "Tableau 5. Gestion des échecs OCR et biométriques", "Tableau 5. Gestion des échecs OCR et biométriques")
    fig3 = find_exact(doc, "Figure 3. MLD simplifié des entités KYC principales")
    cur = fig3
    cur = insert_after(cur, "Figure 4. Cadrage mobile du périmètre documentaire autour des CNI camerounaises", "Front Matter")
    insert_after(cur, "Figure 5. File back office de validation KYC et indicateurs de suivi", "Front Matter")


def replace_bibliography(doc: Document) -> None:
    start = next(i for i, p in enumerate(doc.paragraphs) if p.text.strip() == "Références bibliographiques")
    end = next(i for i, p in enumerate(doc.paragraphs) if i > start and p.text.strip() == "Annexes")
    body = doc.element.body
    for p in list(doc.paragraphs[start + 1 : end]):
        body.remove(p._element)

    entries = [
        "Banque Internationale du Cameroun pour l'Epargne et le Crédit. (2026). A propos de la BICEC. https://www.bicec.com/la-bicec/",
        "Banque Internationale du Cameroun pour l'Epargne et le Crédit. (2026). Réseau BICEC. https://www.bicec.com/la-bicec/reseau/",
        "BICEC VeriPass. (2026). Product Requirements Document interne du projet BICEC VeriPass [Document interne non publié].",
        "BICEC VeriPass. (2026). Backlog GitHub issues et jalons du projet BICEC VeriPass [Dépôt projet interne].",
        "Bulatov, K. B., Bezmaternykh, P. V., Nikolaev, D. P., & Arlazarov, V. V. (2022). Towards a unified framework for identity documents analysis and recognition. Computer Optics, 46(3), 436 454. https://doi.org/10.18287/2412-6179-CO-1024",
        "Carta, S., Giuliani, A., Piano, L., & Tiddia, S. G. (2024). An end to end OCR free solution for identity document information extraction. Procedia Computer Science, 246, 453 462. https://doi.org/10.1016/j.procs.2024.09.425",
        "Celery Project. (2026). Celery documentation. https://docs.celeryq.dev/",
        "Commission Bancaire de l'Afrique Centrale. (2023). Règlement COBAC R 2023/01 relatif aux diligences des établissements assujettis en matière de lutte contre le blanchiment de capitaux, le financement du terrorisme et de la prolifération.",
        "Consultative Group to Assist the Poor. (2021). Regulation for inclusive digital finance. CGAP.",
        "Docker. (2026). Docker Compose documentation. https://docs.docker.com/compose/",
        "Droit Médias Finance. (2024). CEMAC, LCB FT : le règlement COBAC R 2023/01 relatif aux nouvelles diligences anti blanchiment entre en vigueur le 1er juillet 2024. https://www.droitmediasfinance.com/index.php/actualites/droit-bancaire/739-cobac-le-reglement-cobac-r-2023-01-relatif-aux-nouvelles-diligences-des-etablissements-assujettis-en-matiere-de-lutte-contre-le-blanchiment-de-capitaux-entrera-en-vigueur-le-1er-juillet-2024",
        "Du, Y., Li, C., Guo, R., Yin, X., Liu, W., Zhou, J., Bai, Y., Yu, Z., Yang, Y., Dang, Q., & Wang, H. (2020). PP OCR: A practical ultra lightweight OCR system. arXiv. https://arxiv.org/abs/2009.09941",
        "FastAPI. (2026). FastAPI documentation. https://fastapi.tiangolo.com/",
        "Financial Action Task Force. (2020). Guidance on digital identity. FATF. https://www.fatf-gafi.org",
        "OpenAPI Initiative. (2024). OpenAPI Specification. https://spec.openapis.org",
        "PostgreSQL Global Development Group. (2026). PostgreSQL documentation. https://www.postgresql.org/docs/",
        "Puthod, M. (2025). Référentiels de données [Support de cours]. UCAC ICAM.",
        "Redis Ltd. (2026). Redis documentation. https://redis.io/docs/latest/",
        "République du Cameroun. (2024). Loi n° 2024/017 du 23 décembre 2024 relative à la protection des données à caractère personnel au Cameroun. https://www.prc.cm/fr/multimedia/documents/10258-loi-n-2024-017-du-23-12-2024-web?album_id=200",
        "Yu, Z., Qin, Y., Li, X., Zhao, C., Lei, Z., & Zhao, G. (2023). Deep learning for face anti spoofing: A survey. IEEE Transactions on Pattern Analysis and Machine Intelligence, 45(5), 5609 5631.",
    ]
    cur = doc.paragraphs[start]
    for entry in entries:
        cur = insert_after(cur, entry, "Bibliographie")


def guardrails(doc: Document) -> None:
    text = "\n".join(p.text for p in doc.paragraphs)
    forbidden = ["\u2014", "\u2013", "Jean", "Thomas", "Sylvie", "Admin IT"]
    hits = [token for token in forbidden if token in text]
    if hits:
        raise AssertionError(f"Forbidden tokens in document: {hits}")

    spaced_dash = re.findall(r"\s[-]\s", text)
    if spaced_dash:
        raise AssertionError("Spaced dash separators remain in document")


def main() -> None:
    shutil.copyfile(SRC, OUT)
    make_backoffice_redacted()
    doc = Document(OUT)

    replace_contains(
        doc,
        "Conception et développement d'un écosystème intelligent d'acquisition client",
        "Conception et développement d'un écosystème intelligent d'acquisition client\nIntégration de mécanismes avancés de vérification pour l'automatisation de la conformité KYC",
    )

    update_front_lists(doc)
    add_abbreviation_rows(doc)

    replace_exact(
        doc,
        "Les documents de cadrage du projet retiennent un délai d'ouverture pouvant aller de 48 heures à 14 jours dans le processus manuel de la BICEC. Ce délai dépend de la disponibilité du client, de la complétude du dossier, de la lisibilité des pièces et des échanges nécessaires entre agence et back office. Cette friction justifie l'objectif de réduire le temps de constitution du dossier, tout en conservant une revue interne conforme.",
        "Les entretiens de cadrage et le PRD interne BICEC VeriPass retiennent un délai d'ouverture pouvant aller de 48 heures à 14 jours dans le processus manuel de la BICEC (BICEC VeriPass, 2026). Ce délai dépend de la disponibilité du client, de la complétude du dossier, de la lisibilité des pièces et des échanges nécessaires entre agence et back office. Cette friction justifie l'objectif de réduire le temps de constitution du dossier, tout en conservant une revue interne conforme.",
    )

    replace_exact(
        doc,
        "Le règlement COBAC R 2023/01 du 19 décembre 2023 relatif aux diligences des établissements assujettis en matière de lutte contre le blanchiment des capitaux et le financement du terrorisme constitue le cadre juridique de référence du projet dans la zone CEMAC. Il impose aux établissements assujettis d'identifier leurs clients, de vérifier leur identité à partir de sources fiables, de tenir compte du risque et de conserver les éléments utiles à la vigilance. Cette obligation exclut une lecture purement commerciale du KYC et impose de maintenir une décision humaine explicable.",
        "Le règlement COBAC R 2023/01 du 19 décembre 2023 relatif aux diligences des établissements assujettis en matière de lutte contre le blanchiment des capitaux, le financement du terrorisme et la prolifération constitue le cadre juridique de référence du projet dans la zone CEMAC (Commission Bancaire de l'Afrique Centrale, 2023). Il impose aux établissements assujettis d'identifier leurs clients, de vérifier leur identité à partir de sources fiables, de tenir compte du risque et de conserver les éléments utiles à la vigilance. Cette obligation exclut une lecture purement commerciale du KYC et impose de maintenir une décision humaine explicable.",
    )
    insert_paragraphs_after(
        find_contains(doc, "Commission Bancaire de l'Afrique Centrale, 2023"),
        [
            "Les commentaires juridiques disponibles sur ce texte soulignent aussi son entrée en vigueur au 1er juillet 2024 et le renforcement des diligences attendues des acteurs financiers de la CEMAC (Droit Médias Finance, 2024). Pour VeriPass, cette évolution donne un sens concret à l'auditabilité : l'application doit être capable de relier une décision KYC aux pièces, aux contrôles et aux acteurs qui l'ont produite.",
        ],
    )

    replace_exact(
        doc,
        "En parallèle, la Loi n° 2024/017 relative à la protection des données à caractère personnel au Cameroun renforce l'attention portée au consentement, à la minimisation, à la sécurité et à la maîtrise des données sensibles. Pour VeriPass, cette contrainte invite à privilégier une architecture où les pièces d'identité, les données biométriques et les journaux restent sous contrôle de l'organisation.",
        "En parallèle, la Loi n° 2024/017 du 23 décembre 2024 relative à la protection des données à caractère personnel au Cameroun renforce l'attention portée au consentement, à la minimisation, à la sécurité et à la maîtrise des données sensibles (République du Cameroun, 2024). Pour VeriPass, cette contrainte invite à privilégier une architecture où les pièces d'identité, les données biométriques et les journaux restent sous contrôle de l'organisation.",
    )

    insert_paragraphs_after(
        find_exact(doc, "L'impact économique attendu doit donc être formulé avec prudence. VeriPass peut contribuer à réduire le coût d'acquisition client si le parcours diminue les allers retours, accélère la revue des dossiers et réduit les erreurs de saisie. Ces gains devront être mesurés pendant une phase pilote avant toute conclusion chiffrée."),
        [
            "L'évaluation financière doit donc raisonner comme un pilotage de projet. Le coût complet ne se limite pas au serveur ou au développement logiciel : il inclut la machine utilisée, la connexion internet, l'électricité, les consommations d'outils d'assistance et de calcul, le temps humain, les échanges d'encadrement, la documentation et les essais. Cette lecture coût, délai et périmètre permet d'apprécier la valeur d'un prototype qui couvre le parcours mobile, l'API, l'OCR, la biométrie, la base, les workers et le back office sans recourir à une prestation externe complète.",
            "La consommation de tokens et de services d'assistance doit être suivie comme une charge de conception, au même titre que les coûts de connexion ou de calcul. Elle reste maîtrisée lorsque les tâches répétitives sont transformées en scripts, preuves de tests et documents de cadrage réutilisables. Dans une perspective de reproduction industrielle, le juste chiffrage devra distinguer les coûts déjà supportés par le stage, les coûts internes BICEC, et les coûts nouveaux liés à l'hébergement, à la sécurité, au pilote et à l'exploitation.",
        ],
    )

    insert_paragraphs_after(
        find_exact(doc, "La gouvernance des données doit rester proportionnée à la sensibilité du KYC. Les pièces d'identité, les images de visage et les journaux de décision ne sont pas de simples fichiers techniques. Ils portent une finalité réglementaire, doivent être accessibles seulement aux rôles habilités, et doivent rester conservés avec une trace suffisante pour justifier les décisions prises."),
        [
            "La notion de référentiel de données permet de préciser cette exigence. Les données de référence correspondent aux objets stables du métier, tels que l'identité client, le type de document, le statut du dossier, le rôle d'un agent ou le motif d'une décision. Les données transactionnelles décrivent les événements du parcours, comme un upload, une correction OCR, une preuve de vie ou une soumission. Les données décisionnelles agrègent ensuite ces traces pour mesurer la charge, les reprises et les motifs d'anomalie (Puthod, 2025). Cette distinction justifie que VeriPass reste un système OLTP pour l'instruction des dossiers, tout en préparant des indicateurs d'exploitation.",
        ],
    )

    insert_paragraphs_after(
        find_exact(doc, "Les recherches récentes sur l'extraction d'informations dans les documents d'identité montrent aussi la rareté des jeux de données réellement exploitables, pour des raisons de confidentialité et de sécurité (Carta et al., 2024). Cette contrainte justifie, dans le cadre du mémoire, de distinguer ce qui est implémenté, ce qui est mesuré sur les preuves disponibles, et ce qui devra être calibré lors d'un pilote plus large."),
        [
            "Cette rareté a eu une conséquence directe sur le projet. Il n'était pas acceptable de constituer librement un corpus de CNI réelles, car ces pièces exposent des identités, des signatures, des photos et des numéros sensibles. Le travail expérimental a donc combiné images disponibles dans un cadre contrôlé, génération exploratoire, inspection visuelle, et transformation des erreurs observées en règles générales de post traitement. Les notebooks Marimo ont joué le rôle de laboratoire : ils permettaient de lancer une expérience OCR, d'observer les blocs reconnus, de comparer les sorties, puis de rapatrier seulement les règles validées dans le service backend.",
            "La spécificité des CNI camerounaises anciennes a également pesé sur les choix. Les reflets du support, l'usure, les variations de cadrage, les libellés bilingues, les professions, les postes d'identification et les adresses rendent insuffisante une simple extraction texte. Le pipeline doit reconnaître un champ dans son contexte, distinguer recto et verso, préserver des valeurs rares mais légitimes, et ne pas rejeter silencieusement une information parce qu'elle ressemble à un bruit OCR.",
        ],
    )

    insert_paragraphs_after(
        find_exact(doc, "La performance est donc traitée comme une contrainte de parcours, non comme une promesse isolée. Les traitements longs doivent être asynchrones, les statuts doivent rester lisibles pour l'utilisateur, et les erreurs doivent ouvrir une reprise contrôlée. La sécurité, de son côté, repose sur la séparation des rôles, l'authentification, la journalisation, l'intégrité documentaire et la limitation des accès aux données sensibles."),
        [
            "Une autre contrainte non fonctionnelle a été la sobriété de calcul. Le prototype devait rester exécutable sur une machine de développement de gamme moyenne disposant de 16 Go de RAM. Cette contrainte a imposé de charger les modèles avec prudence, de limiter la concurrence des traitements lourds, de séparer les tâches OCR et GLM, et de considérer la latence comme un phénomène d'exploitation autant que comme une métrique algorithmique.",
        ],
    )

    insert_paragraphs_after(
        find_exact(doc, "Cette approche permet d'intégrer les retours d'encadrement sans perdre la maîtrise des exigences. Les livrables techniques (contrats API FastAPI, modèles relationnels PostgreSQL, PWA React/Vite, back office, journaux de sécurité et preuves d'exécution) sont produits par incréments, testés localement et documentés afin de préparer la démonstration finale."),
        [
            "La phase de conception a occupé une part importante du stage. Les documents issus du cadrage, notamment charte projet, PRD, backlog, spécifications UX, décisions d'architecture, revues de risques et supports de présentation, ont servi de supports de discussion avec des profils différents. Ils ne doivent pas être lus comme de simples annexes techniques : ils matérialisent la capacité du projet à dialoguer avec le métier, la conformité, l'exploitation, l'expérience utilisateur et le pilotage.",
            "Cette organisation montre aussi une progression managériale. Le projet a demandé de hiérarchiser les risques, de défendre des arbitrages, de maintenir un fil de décision, et d'accepter que certaines ambitions soient reportées pour préserver la cohérence du périmètre. Dans un environnement bancaire, cette discipline compte autant que la production de code, car une solution non cadrée devient rapidement difficile à valider.",
        ],
    )

    insert_paragraphs_after(
        find_exact(doc, "Ce cadre méthodologique s'incarne dans un calendrier structuré par jalons. Les risques suivis concernent principalement la qualité des images, les performances OCR, la robustesse biométrique, la sécurité des pièces et la validation métier des parcours. Les échanges avec M. Jackson KOMBE LELE et les équipes concernées permettent d'ajuster les compromis de conception au fur et à mesure."),
        [
            "Le suivi opérationnel du dépôt confirme cette structuration. Les métadonnées GitHub consultées le 2 juin 2026 font apparaître 197 issues visibles, dont 30 closes, réparties sur des jalons fonctionnels comme M0 Dev Ready, M1 Auth Complete, M2 Capture MVP, M3 AI Engine Complete, M4 Back Office MVP, M5 Full KYC Flow et M6 Analytics and Demo (BICEC VeriPass, 2026). La répartition par labels, notamment mobile, backend, back office, KYC, OCR, biométrie, sécurité et conformité, montre que le projet n'a pas été conduit comme une suite de tâches isolées, mais comme un portefeuille de travaux coordonnés.",
            "Cette granularité a servi la constance du travail. Les jalons ont permis de séparer les fondations d'authentification, la capture, le moteur IA, la revue back office, le parcours complet et les preuves de démonstration. Elle a aussi permis de documenter les arbitrages lorsque la conformité, la performance ou la disponibilité machine imposaient de modifier l'ordre des priorités.",
        ],
    )

    insert_paragraphs_after(
        find_exact(doc, "Redis et Celery forment la couche de traitement différé. Elle devient nécessaire dès que le temps de traitement dépasse l'attente normale d'un appel HTTP. OCR, notifications, tâches planifiées et certains contrôles de cohérence peuvent ainsi être exécutés sans bloquer le client. Le statut du dossier devient alors le moyen de synchroniser l'expérience utilisateur avec les traitements en arrière plan."),
        [
            "Cette architecture répond aussi à une contrainte matérielle précise. Les notes d'analyse WSL2 estiment que la pile sans modèles chargés reste autour de quelques gigaoctets, tandis que l'API avec PaddleOCR, le worker GLM et les autres services peuvent approcher 8,5 Go. Sur une machine de 16 Go de RAM, l'objectif n'est donc pas de lancer toutes les charges en concurrence maximale, mais de contrôler l'activation des modèles, les files, les verrous et les reprises.",
            "Redis n'est pas choisi comme base métier. Il joue le rôle de broker, de cache court et de support de verrous, tandis que PostgreSQL reste la source de vérité. Les verrous de type OCR, GLM et agent permettent d'éviter les traitements simultanés dangereux, notamment lorsqu'un moteur IA consomme beaucoup de mémoire ou lorsqu'un dossier doit être affecté à un agent selon la charge active.",
        ],
    )

    insert_paragraphs_after(
        find_exact(doc, "Lorsqu'un document est reçu, il n'est pas seulement enregistré comme fichier. Le système calcule un hash, conserve le chemin relatif, le type de document, la date de capture, la taille et le statut OCR. Cette métadonnée permet de vérifier l'intégrité du fichier et de relier la pièce à la session correcte."),
        [
            "Le traitement documentaire prend aussi en compte le coût réel des images. Une capture mobile peut atteindre une taille inutilement élevée pour l'OCR, surtout si le document occupe seulement une partie de l'image. Le projet prévoit donc une compression contrôlée et une normalisation avant certains traitements afin de réduire la charge mémoire et la latence, sans dégrader les zones utiles à la lecture. Ce compromis est typiquement un arbitrage d'ingénieur : une image trop lourde ralentit le pipeline, mais une compression excessive détruit les petits caractères.",
            "La sécurité du flux documentaire s'appuie sur plusieurs niveaux. Le hash SHA 256 assure l'intégrité, les rôles limitent l'accès, et les chemins documentaires sont séparés de la base relationnelle. Le code comporte également des fonctions de compression, chiffrement et déchiffrement AES GCM pour les flux OCR externes ou sensibles, tandis que la production devra formaliser la politique de clés, de sauvegardes chiffrées et de conservation.",
        ],
    )

    insert_paragraphs_after(
        find_exact(doc, "PaddleOCR est adapté à une première passe, car il permet une extraction rapide et locale. Dans le contexte VeriPass, cette rapidité est utile pour donner un retour immédiat ou quasi immédiat sur une capture. Elle ne dispense pas d'une seconde lecture lorsque le document est bruité, lorsque des champs sont absents ou lorsque le format réel diffère des hypothèses prévues."),
        [
            "L'amélioration de l'OCR s'est faite par itérations. La boucle de test du 26 mai 2026 a montré une première passe complète de 53 succès sur 60 identifiants CNI, puis des reprises ciblées qui ont porté l'agrégat final à 60 identifiants avec au moins une passe réelle réussie dans le parcours applicatif. Les corrections ont notamment porté sur les dates à dix ans, les adresses verso, les professions comme MENAGERE ou INGENIEUR, les numéros CNI et les postes d'identification. Ce résultat n'est pas une précision statistique définitive, mais il prouve que l'OCR a été travaillé comme un système calibré, testé et corrigé.",
        ],
    )

    insert_paragraphs_after(
        find_exact(doc, "PostgreSQL est également pertinent parce que les écritures KYC exigent des transactions. Lorsqu'un dossier est soumis, l'état, l'adresse IP, les documents, les consentements et les notifications doivent rester cohérents. Une base relationnelle permet de contrôler ces liens et de limiter les anomalies de cycle de vie."),
        [
            "Le choix PostgreSQL se justifie aussi par la distinction entre OLTP et OLAP. VeriPass doit d'abord instruire des dossiers, enregistrer des décisions et préserver des transitions cohérentes ; il relève donc d'un besoin transactionnel. Une architecture OLAP serait pertinente pour analyser les volumes, les délais et les motifs de rejet à partir de données agrégées, mais elle ne doit pas porter la décision KYC en temps réel. Le modèle retenu prépare donc le décisionnel sans sacrifier la cohérence opérationnelle.",
            "Dans cette logique, le référentiel de données n'est pas un entrepôt analytique. Il identifie les objets métier stables, les entités persistantes et les relations qui rendent le dossier interprétable. Les champs JSONB absorbent les sorties variables des moteurs, mais les responsabilités critiques restent relationnelles : dossier, document, champ, biométrie, décision et audit.",
        ],
    )

    insert_paragraphs_after(
        find_exact(doc, "Le choix d'un environnement conteneurisé répond aussi à une exigence pédagogique et industrielle. Il permet de montrer l'ensemble de la chaîne sans dépendre d'un service externe pour les données sensibles. Les volumes documentaires, les volumes de base et les volumes de sauvegarde matérialisent la séparation entre code applicatif, métadonnées relationnelles et fichiers KYC."),
        [
            "Le choix de Docker n'est donc pas un simple confort de développement. Il permet de reproduire la pile, d'isoler les services, de limiter les écarts entre démonstration et reprise technique, et de rendre visibles les dépendances critiques. Dans un projet où OCR, base, broker, workers, interfaces et proxy doivent fonctionner ensemble, la conteneurisation devient un outil de maîtrise du système plutôt qu'un emballage.",
        ],
    )

    cur = insert_figure_after(
        find_exact(doc, "Le mobile ne prend pas la décision KYC. Il collecte les preuves et rend les corrections possibles. Cette limite est importante : le client peut confirmer ou corriger des champs OCR, mais la décision bancaire reste dans le back office, avec un agent habilité et une trace exploitable."),
        MOBILE_SCOPE,
        "Figure 4. Cadrage mobile du périmètre documentaire autour des CNI camerounaises",
        "Source : capture d'écran du prototype VeriPass, preuve de test KYC conformité.",
        3.15,
    )
    insert_paragraphs_after(
        cur,
        [
            "La figure 4 illustre un choix volontaire de périmètre : le MVP n'ouvre pas toutes les catégories de pièces, il concentre le flux sur la CNI camerounaise. Cette restriction réduit la dispersion fonctionnelle et permet de travailler plus finement les contraintes réelles du document retenu.",
        ],
    )

    insert_paragraphs_after(
        find_exact(doc, "PaddleOCR sert de première passe locale. Son intérêt est de fournir un résultat rapide sans envoyer les pièces d'identité vers un service externe. Le pipeline applique ensuite des règles de post traitement adaptées aux champs attendus d'une CNI : noms, prénoms, dates, numéro de document, poste d'identification, profession ou adresse selon le côté de la pièce et la lisibilité de la capture."),
        [
            "Le post traitement est devenu une partie centrale de la réalisation. Il ne s'agit pas seulement de nettoyer un texte : il faut interpréter la zone du document, protéger les champs légitimes, refuser les valeurs manifestement hors contexte, normaliser les dates et éviter qu'une valeur du verso soit utilisée comme donnée du recto. Cette couche a beaucoup compté dans l'amélioration des résultats sur les CNI anciennes.",
        ],
    )

    insert_paragraphs_after(
        find_exact(doc, "La correction humaine est intégrée au pipeline. Le client peut confirmer ou corriger les champs dans le parcours, et le back office peut intervenir sur les champs lors de la revue. La valeur initiale, la valeur corrigée, l'auteur et l'horodatage doivent rester disponibles pour comprendre l'écart entre l'extraction automatique et la décision finale."),
        [
            "Les notebooks d'expérimentation ont joué ici un rôle de passerelle entre recherche et application. Ils ont permis de lancer des traitements exploratoires, de visualiser les blocs, de comparer les moteurs, puis de transformer les observations en règles backend. Ce fonctionnement évite d'introduire trop tôt une hypothèse fragile dans l'application, tout en accélérant l'apprentissage sur des cas difficiles.",
        ],
    )

    cur = insert_figure_after(
        find_exact(doc, "L'audit log complète cette logique. Il donne une trace transversale des actions sensibles : revue, affectation, classification, correction, décision, sauvegarde ou export. Cette trace ne remplace pas la conformité métier, mais elle fournit un support technique à la reconstitution des dossiers."),
        BACKOFFICE_REDACTED,
        "Figure 5. File back office de validation KYC et indicateurs de suivi",
        "Source : capture d'écran anonymisée du prototype VeriPass.",
        6.1,
    )
    insert_paragraphs_after(
        cur,
        [
            "La figure 5 montre que le back office n'est pas une simple liste de dossiers. Il doit aider l'agent à prioriser, comprendre les anomalies, mesurer la charge et prendre une décision motivée. Cette interface est donc un instrument d'organisation autant qu'une interface technique.",
        ],
    )

    insert_paragraphs_after(
        find_exact(doc, "La deuxième difficulté concerne les traitements longs. L'OCR et le liveness peuvent dépasser le temps acceptable d'un appel interactif. L'usage de Redis et Celery répond à cette contrainte en séparant la requête client du traitement différé. Le client ne doit pas attendre sans explication; le dossier doit porter un statut lisible."),
        [
            "Le parallélisme a dû être conçu de manière sélective. Certaines tâches peuvent être traitées en file sans risque, comme les notifications ou des traitements légers. D'autres, notamment GLM OCR et certains contrôles d'image, doivent être limités par des verrous et par une affectation selon la charge active. Le principe retenu rapproche l'affectation de la logique Weighted Round Robin et Least Connections : répartir les dossiers quand cela est possible, mais bloquer la concurrence lorsque la mémoire ou la cohérence du dossier l'exige.",
        ],
    )

    insert_paragraphs_after(
        find_exact(doc, "Le dernier arbitrage concerne l'intégration bancaire. Le prototype n'intègre pas directement un core banking de production. Ce choix évite de confondre preuve de concept KYC et activation commerciale réelle. Il permet de concentrer le mémoire sur la chaîne d'acquisition, de vérification, de traçabilité et de revue."),
        [
            "Ces arbitrages montrent que l'apport ne se limite pas à l'assemblage d'outils. Faire fonctionner OCR, liveness, back office, base transactionnelle, cache, workers et documents sur une machine moyenne gamme a imposé de raisonner sur la mémoire, les files, les verrous, les seuils, les reprises, les secrets et la preuve. C'est cette maîtrise des contraintes qui distingue un prototype d'ingénierie d'une simple démonstration technique.",
        ],
    )

    replace_exact(
        doc,
        "La boucle OCR documentée dans le dépôt a porté sur un lot de 60 identifiants CNI. Le rapport indique qu'au terme des passes initiales et des reprises ciblées, chaque identifiant CNI a obtenu au moins une passe OCR réelle réussie dans le parcours applicatif. Ce résultat est utile, car il montre que le pipeline fonctionne sur un ensemble varié de captures, avec interaction réelle avec les endpoints et les données de revue.",
        "La boucle OCR documentée dans le dépôt a porté sur un lot de 60 identifiants CNI. Le rapport du 26 mai 2026 indique une première passe complète à 53 succès sur 60, puis des reprises ciblées ayant abouti à un agrégat final de 60 identifiants avec au moins une passe OCR réelle réussie dans le parcours applicatif. Ce résultat est utile, car il montre que le pipeline fonctionne sur un ensemble varié de captures, avec interaction réelle avec les endpoints et les données de revue.",
    )
    insert_paragraphs_after(
        find_exact(doc, "Les corrections apportées dans la boucle OCR sont instructives. Certaines concernaient l'adresse, la profession, les dates, le numéro CNI ou le poste d'identification. Elles montrent que le travail d'ingénierie ne se limite pas à appeler un moteur OCR : il faut normaliser les formats, protéger les valeurs légitimes, tenir compte des zones du document et conserver l'incertitude lorsque la capture reste ambiguë."),
        [
            "La vérification directe des fonctions de post traitement a aussi donné 8 tests réussis sur 8. Cette preuve reste ciblée, mais elle renforce la crédibilité du pipeline : les règles n'ont pas seulement été écrites après observation, elles ont été vérifiées sur les cas précis qui avaient provoqué les erreurs.",
        ],
    )

    insert_paragraphs_after(
        find_exact(doc, "Les workers constituent un point sensible de l'exploitation. Si la file OCR ralentit, l'application peut rester accessible tout en bloquant la progression des dossiers. La supervision doit donc distinguer disponibilité web, disponibilité de la base, capacité des files et état des traitements différés. Cette distinction est indispensable pour diagnostiquer rapidement une panne partielle."),
        [
            "La contrainte mémoire renforce cette exigence. En environnement local, la charge estimée avec PaddleOCR, GLM OCR et les services associés peut approcher 8,5 Go. La supervision doit donc observer non seulement les erreurs fonctionnelles, mais aussi les consommations mémoire, la longueur des files, les tâches en attente et les verrous expirés. Un système KYC peut sembler disponible tout en étant incapable de traiter les dossiers si ses workers lourds sont saturés.",
        ],
    )

    insert_paragraphs_after(
        find_exact(doc, "L'impact humain est également important. Le projet montre que l'automatisation bancaire ne consiste pas à remplacer les acteurs, mais à déplacer leur effort vers les tâches à plus forte valeur : arbitrer, contrôler, expliquer, traiter les exceptions et améliorer le dispositif."),
        [
            "Sur le plan organisationnel, le stage a aussi renforcé les compétences de coordination. Les échanges avec l'encadrement, la tenue d'un backlog, les documents de cadrage, les preuves de tests et les arbitrages techniques ont imposé une posture de chef de projet autant que de développeur. Il a fallu expliquer les choix, écouter les contraintes métier, accepter les procédures, puis transformer ces contraintes en décisions vérifiables.",
            "Sur le plan financier, l'impact devra être évalué à partir d'indicateurs de pilote : temps moyen de constitution, taux de reprises, temps agent par dossier, coût de stockage, consommation IA, coût de connexion, énergie et maintenance. Le mémoire ne fixe pas un retour sur investissement artificiel, mais il fournit une méthode pour comparer le coût de reproduction du prototype avec une prestation équivalente du marché sur un périmètre comparable.",
        ],
    )

    insert_paragraphs_after(
        find_exact(doc, "La perspective la plus immédiate consiste à transformer les preuves techniques en indicateurs de pilotage. Il faudra mesurer le temps moyen de constitution, le taux de dossiers incomplets, le nombre de corrections OCR, les motifs de demande de complément, les cas de rejet biométrique et la charge par rôle back office. Ces indicateurs permettront de juger la solution sur son utilité réelle, et non sur la seule existence des composants techniques."),
        [
            "Une perspective complémentaire concerne la gouvernance des coûts. La banque devra décider quels traitements restent locaux, lesquels peuvent être externalisés sous conditions, et comment suivre les consommations de calcul, de stockage et d'assistance IA. Cette gouvernance est indispensable pour que l'innovation reste économiquement soutenable et compatible avec les exigences de souveraineté des données.",
        ],
    )

    replace_exact(
        doc,
        "Sur le plan technique, le projet a permis de mettre en pratique des compétences d'ingénierie logiciel et data. La conception a exigé de modéliser les entités KYC, de définir des états de dossier, de relier les documents à leurs métadonnées, de traiter les sorties OCR comme des données incertaines et de conserver les traces nécessaires à l'audit. La réalisation a également renforcé la maîtrise des contrats API, de l'orchestration asynchrone, des conteneurs Docker, des files Celery et de la séparation entre collecte, traitement et décision.",
        "Sur le plan technique, le projet a permis de mettre en pratique des compétences d'ingénierie logiciel, data et IA appliquée. La conception a exigé de modéliser les entités KYC, de distinguer référentiel, données transactionnelles et données décisionnelles, de traiter les sorties OCR comme des données incertaines, puis de relier documents, scores, corrections, décisions et journaux. La réalisation a renforcé la maîtrise des contrats API, de l'orchestration asynchrone, des conteneurs Docker, des files Celery, de Redis, de PostgreSQL, des seuils, des verrous, de la compression, du chiffrement et de la séparation entre collecte, traitement et décision.",
    )
    insert_paragraphs_after(
        find_exact(doc, "Sur le plan organisationnel, le stage a montré que la qualité d'une solution bancaire ne dépend pas seulement de sa performance technique. Elle dépend aussi de sa capacité à s'intégrer aux rôles existants, à respecter les responsabilités des agents, à rendre les anomalies compréhensibles et à produire des preuves exploitables. Les échanges avec l'encadrement BICEC ont donc orienté le projet vers une automatisation assistée, et non vers une décision automatique."),
        [
            "Sur le plan humain, le projet a demandé une posture de leadership mesurée. Il a fallu organiser le travail, construire des documents de dialogue, maintenir le cap malgré les procédures, transformer les critiques en améliorations, et tenir une trajectoire cohérente jusqu'au prototype. Cette dimension est essentielle dans un mémoire d'ingénieur, car la qualité finale d'un système bancaire dépend aussi de la manière dont les acteurs comprennent et acceptent ses choix.",
            "Sur le plan économique, la solution doit encore être évaluée par un pilote, mais le mémoire fournit déjà une grille de coût. Elle couvre le temps de conception, les ressources humaines, la machine, la connexion, l'électricité, les outils d'assistance, les consommations IA, l'hébergement, la supervision, le stockage, la sécurité et la maintenance. Cette grille permettra de comparer l'effort du stage avec le coût d'une réalisation équivalente sur le marché, sans confondre prototype académique et produit industriel.",
        ],
    )

    replace_bibliography(doc)
    guardrails(doc)
    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
