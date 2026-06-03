from __future__ import annotations

import shutil
from datetime import date, timedelta
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Inches
from docx.text.paragraph import Paragraph
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[2]
SRC = ROOT / "docs" / "rapport-stage" / "memoire-stage-andre-yoann-kenmogne-v9.docx"
OUT = ROOT / "docs" / "rapport-stage" / "memoire-stage-andre-yoann-kenmogne-v10.docx"
FIG_DIR = ROOT / "docs" / "rapport-stage" / "figures"
FIG_TECH = FIG_DIR / "v10_budget_ram_ocr.png"
FIG_COST = FIG_DIR / "v10_cout_transport.png"


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    candidates = [
        Path("C:/Windows/Fonts/arialbd.ttf") if bold else Path("C:/Windows/Fonts/arial.ttf"),
        Path("C:/Windows/Fonts/calibrib.ttf") if bold else Path("C:/Windows/Fonts/calibri.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size=size)
    return ImageFont.load_default()


def text_center(draw: ImageDraw.ImageDraw, xy: tuple[int, int], text: str, fnt, fill=(30, 30, 30)) -> None:
    bbox = draw.textbbox((0, 0), text, font=fnt)
    draw.text((xy[0] - (bbox[2] - bbox[0]) // 2, xy[1]), text, font=fnt, fill=fill)


def draw_technical_figure() -> None:
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    img = Image.new("RGB", (1500, 760), "white")
    d = ImageDraw.Draw(img)
    title = font(30, True)
    h2 = font(22, True)
    body = font(18)
    small = font(15)
    accent = (231, 112, 0)
    blue = (32, 92, 137)
    gray = (230, 235, 240)
    dark = (40, 45, 52)

    text_center(d, (750, 22), "Budget technique : mémoire, seuils et amélioration OCR", title, dark)
    d.line((70, 70, 1430, 70), fill=(210, 210, 210), width=2)

    d.text((95, 95), "Budget RAM sur poste 16 Go", font=h2, fill=dark)
    d.text((845, 95), "Boucle OCR sur CNI de test", font=h2, fill=dark)

    values = [("Pile de base", 2.5, blue), ("PaddleOCR", 2.5, accent), ("GLM OCR", 3.5, (98, 126, 170))]
    x0, y0 = 110, 620
    scale = 34
    x = x0
    total = 0
    for label, val, color in values:
        h = int(val * scale)
        d.rectangle((x, y0 - h, x + 120, y0), fill=color)
        d.text((x, y0 + 14), label, font=small, fill=dark)
        d.text((x + 28, y0 - h - 28), f"{val:.1f} Go".replace(".", ","), font=small, fill=dark)
        total += val
        x += 150
    d.line((90, y0 - int(16 * scale), 610, y0 - int(16 * scale)), fill=(180, 40, 40), width=3)
    d.text((620, y0 - int(16 * scale) - 13), "RAM machine : 16 Go", font=small, fill=(180, 40, 40))
    d.text((110, 155), "Equation de charge :", font=body, fill=dark)
    d.text((110, 190), "M_total = M_base + M_Paddle + M_GLM", font=body, fill=dark)
    d.text((110, 225), "M_total = 2,5 + 2,5 + 3,5 = 8,5 Go", font=body, fill=dark)
    d.text((110, 260), "Principe : charger, verrouiller, traiter, libérer.", font=body, fill=dark)
    d.rectangle((95, 300, 680, 335), fill=gray)
    d.text((110, 306), "Redis garde les files et verrous ; PostgreSQL garde la vérité métier.", font=small, fill=dark)

    # OCR bars
    labels = [("Passe initiale", 53, 60, blue), ("Après reprises", 60, 60, accent), ("Tests règles", 8, 8, (70, 150, 100))]
    bx, by = 860, 620
    maxv = 60
    for label, val, denom, color in labels:
        h = int((val / maxv) * 390)
        d.rectangle((bx, by - h, bx + 145, by), fill=color)
        d.text((bx, by + 14), label, font=small, fill=dark)
        d.text((bx + 42, by - h - 28), f"{val}/{denom}", font=small, fill=dark)
        bx += 170
    d.text((845, 155), "Equation de gating OCR :", font=body, fill=dark)
    d.text((845, 190), "S_ocr(d) = somme(w_f c_f) / somme(w_f)", font=body, fill=dark)
    d.text((845, 225), "GLM(d) = 1 si min(c_f) < 0,90 ou n_champs < 2", font=body, fill=dark)
    d.text((845, 260), "Edition(d) = 1 si S_ocr(d) < 0,95", font=body, fill=dark)
    d.rectangle((845, 300, 1425, 335), fill=gray)
    d.text((860, 306), "Le 60/60 signifie au moins une revue exploitable, pas une précision parfaite.", font=small, fill=dark)

    img.save(FIG_TECH)


def weekday_count(start: date, end: date) -> int:
    days = 0
    cur = start
    while cur <= end:
        if cur.weekday() < 5:
            days += 1
        cur += timedelta(days=1)
    return days


def draw_cost_figure() -> None:
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    start = date(2026, 1, 19)
    end = date(2026, 6, 19)
    daily = 2200
    month_labels: list[str] = []
    month_values: list[int] = []
    cur = date(2026, 1, 31)
    while cur <= end:
        actual = min(cur, end)
        month_labels.append(actual.strftime("%d/%m"))
        month_values.append(weekday_count(start, actual) * daily)
        if cur.month == 12:
            cur = date(cur.year + 1, 1, 31)
        else:
            next_month = cur.month + 1
            next_year = cur.year
            last = date(next_year, next_month + 1, 1) - timedelta(days=1) if next_month < 12 else date(next_year, 12, 31)
            cur = last
    if month_labels[-1] != end.strftime("%d/%m"):
        month_labels.append(end.strftime("%d/%m"))
        month_values.append(weekday_count(start, end) * daily)

    img = Image.new("RGB", (1500, 560), "white")
    d = ImageDraw.Draw(img)
    title = font(30, True)
    body = font(18)
    small = font(15)
    dark = (40, 45, 52)
    accent = (231, 112, 0)
    blue = (32, 92, 137)

    text_center(d, (750, 22), "Coût transport cumulé du stage, hypothèse 5 jours par semaine", title, dark)
    d.text((85, 82), "Période : 19/01/2026 au 19/06/2026 inclus, hors jours fériés non déduits.", font=body, fill=dark)
    d.text((85, 112), "Coût aller retour journalier : 2 200 FCFA. Nombre de jours ouvrés : 110.", font=body, fill=dark)

    x0, y0 = 120, 470
    w, h = 1250, 285
    d.line((x0, y0, x0 + w, y0), fill=dark, width=2)
    d.line((x0, y0, x0, y0 - h), fill=dark, width=2)
    maxv = max(month_values)
    prev = None
    for i, (label, val) in enumerate(zip(month_labels, month_values)):
        x = x0 + int(i * w / (len(month_values) - 1))
        y = y0 - int((val / maxv) * h)
        d.ellipse((x - 6, y - 6, x + 6, y + 6), fill=accent)
        d.text((x - 28, y0 + 12), label, font=small, fill=dark)
        if prev:
            d.line((prev[0], prev[1], x, y), fill=blue, width=4)
        prev = (x, y)
    d.text((x0 + w - 230, y0 - h - 25), f"Total : {maxv:,} FCFA".replace(",", " "), font=body, fill=accent)
    d.text((85, 520), "Equation : C_transport = 110 x 2 200 = 242 000 FCFA.", font=body, fill=dark)
    img.save(FIG_COST)


def set_text(paragraph: Paragraph, text: str) -> None:
    style = paragraph.style
    for run in paragraph.runs:
        run.text = ""
    paragraph.add_run(text)
    paragraph.style = style


def insert_after(paragraph: Paragraph, text: str = "", style: str | None = None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    p = Paragraph(new_p, paragraph._parent)
    if style:
        p.style = style
    if text:
        p.add_run(text)
    return p


def insert_paragraphs_after(paragraph: Paragraph, paragraphs: list[str], style: str = "Memoire Body") -> Paragraph:
    cur = paragraph
    for text in paragraphs:
        cur = insert_after(cur, text, style)
    return cur


def insert_equation_after(paragraph: Paragraph, equation: str) -> Paragraph:
    p = insert_after(paragraph, equation, "Memoire Body")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs:
        run.italic = True
    return p


def insert_figure_after(paragraph: Paragraph, image_path: Path, caption: str, source: str, width: float) -> Paragraph:
    p = insert_after(paragraph, "", "Memoire Body")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(image_path), width=Inches(width))
    cap = insert_after(p, caption, "Caption Memoire")
    src = insert_after(cap, source, "Source Memoire")
    return src


def find_exact(doc: Document, text: str) -> Paragraph:
    for p in doc.paragraphs:
        if p.text.strip() == text:
            return p
    raise ValueError(f"Paragraph not found: {text!r}")


def find_contains(doc: Document, needle: str) -> Paragraph:
    for p in doc.paragraphs:
        if needle in p.text:
            return p
    raise ValueError(f"Paragraph containing {needle!r} not found")


def replace_exact(doc: Document, old: str, new: str) -> None:
    set_text(find_exact(doc, old), new)


def replace_all(doc: Document, old: str, new: str) -> None:
    for p in doc.paragraphs:
        if old in p.text:
            set_text(p, p.text.replace(old, new))


def update_front_lists(doc: Document) -> None:
    fig5 = find_exact(doc, "Figure 5. File back office de validation KYC et indicateurs de suivi")
    cur = insert_after(fig5, "Figure 6. Budget mémoire, logique de seuils et amélioration OCR", "Front Matter")
    insert_after(cur, "Figure 7. Coût transport cumulé du stage selon l'hypothèse cinq jours par semaine", "Front Matter")


def add_abbreviation_rows(doc: Document) -> None:
    table = None
    for t in doc.tables:
        if t.rows and [c.text.strip() for c in t.rows[0].cells][:2] == ["Abréviation", "Signification"]:
            table = t
            break
    if table is None:
        return
    existing = {row.cells[0].text.strip() for row in table.rows[1:]}
    additions = [
        ("AES", "Advanced Encryption Standard, algorithme de chiffrement symétrique"),
        ("HTTP", "HyperText Transfer Protocol, protocole d'échange web"),
        ("IP", "Internet Protocol, adresse réseau utile à la traçabilité technique"),
        ("MLD", "Modèle logique de données"),
        ("MVP", "Minimum Viable Product, périmètre minimal exploitable pour démonstration"),
        ("OTP", "One Time Password, code à usage unique"),
        ("TLS", "Transport Layer Security, protocole de chiffrement des échanges"),
    ]
    for key, value in additions:
        if key in existing:
            continue
        row = table.add_row()
        row.cells[0].text = key
        row.cells[1].text = value


def add_bibliography_entries(doc: Document) -> None:
    refs = find_exact(doc, "Références bibliographiques")
    # Insert after the heading so Word keeps the bibliography in one block. Order is less important than presence.
    entries = [
        "Cui, C., Zhang, Y., Sun, T., Wang, X., Liu, H., Lin, M., et al. (2026). PP OCRv5: A specialized 5M parameter model rivaling billion parameter vision language models on OCR tasks. arXiv. https://arxiv.org/abs/2603.24373",
        "Duan, S., Xue, Y., Wang, W., Su, Z., Liu, H., Yang, S., et al. (2026). GLM OCR technical report. arXiv. https://arxiv.org/abs/2603.10910",
        "PaddleOCR. (2026). OCR pipeline usage documentation. https://www.paddleocr.ai/main/en/version3.x/pipeline_usage/OCR.html",
        "zai org. (2026). GLM OCR: Accurate, fast, comprehensive. GitHub. https://github.com/zai-org/GLM-OCR",
    ]
    cur = refs
    existing = "\n".join(p.text for p in doc.paragraphs)
    for entry in entries:
        if entry[:35] in existing:
            continue
        cur = insert_after(cur, entry, "Bibliographie")


def final_checks(doc: Document) -> None:
    text = "\n".join(p.text for p in doc.paragraphs)
    forbidden = ["\u2014", "\u2013", "Jean", "Thomas", "Sylvie", "Admin IT", "15 minutes", "800 collaborateurs"]
    hits = [token for token in forbidden if token in text]
    if hits:
        raise AssertionError(f"Forbidden tokens: {hits}")
    if " - " in text:
        raise AssertionError("Spaced dash separator found")


def main() -> None:
    draw_technical_figure()
    draw_cost_figure()
    shutil.copyfile(SRC, OUT)
    doc = Document(OUT)

    replace_all(doc, "GLM-OCR", "GLM OCR")
    replace_all(doc, "GLM/OCR", "GLM OCR")
    replace_all(doc, "PP-OCR", "PP OCR")

    update_front_lists(doc)
    add_abbreviation_rows(doc)

    replace_exact(
        doc,
        "Dans ce contexte, l'approche retenue par VeriPass reste volontairement hybride. PaddleOCR est utilisé comme première passe locale pour détecter et reconnaître le texte, dans l'esprit des systèmes OCR légers conçus pour équilibrer vitesse et précision (Du et al., 2020). Lorsque la qualité de capture ou la structure du document rend cette première passe insuffisante, le pipeline prévoit un traitement différé et une revue humaine, plutôt qu'une acceptation silencieuse d'un champ incertain.",
        "Dans ce contexte, l'approche retenue par VeriPass reste hybride et graduée. PaddleOCR est utilisé comme première passe locale parce que la famille PP OCR a été conçue comme une chaîne spécialisée, légère et rapide, où la détection et la reconnaissance sont séparées pour réduire la charge de calcul et les hallucinations propres aux modèles visuels généralistes (Du et al., 2020 ; Cui et al., 2026 ; PaddleOCR, 2026). Lorsque la qualité de capture ou la structure du document rend cette première passe insuffisante, le pipeline prévoit un traitement différé et une revue humaine, plutôt qu'une acceptation silencieuse d'un champ incertain.",
    )

    insert_paragraphs_after(
        find_exact(doc, "Pour la vérification d'identité faciale, la reconnaissance du visage ne suffit pas. Les attaques par présentation, notamment photo, vidéo ou masque, obligent à combiner comparaison faciale et détection de vivacité. Les revues récentes sur le face anti spoofing soulignent que les méthodes doivent être évaluées face à des conditions de capture et à des attaques variées (Yu et al., 2023). Dans VeriPass, cette orientation se traduit par un challenge actif guidé par l'interface, puis par une comparaison entre le selfie et la photo de la CNI. Les résultats restent exprimés en scores, seuils et motifs afin de permettre une revue humaine en cas d'incertitude."),
        [
            "La conclusion scientifique de cette section est donc limitée mais utile : la littérature ne valide pas directement VeriPass sur les CNI camerounaises, elle justifie plutôt une stratégie d'ingénierie prudente. Cette stratégie combine un moteur OCR local pour le débit, un renfort multimodal pour les cas ambigus, des seuils explicites, une correction humaine et une conservation des preuves nécessaires à l'audit.",
        ],
    )

    insert_paragraphs_after(
        find_exact(doc, "Cette architecture répond aussi à une contrainte matérielle précise. Les notes d'analyse WSL2 estiment que la pile sans modèles chargés reste autour de quelques gigaoctets, tandis que l'API avec PaddleOCR, le worker GLM et les autres services peuvent approcher 8,5 Go. Sur une machine de 16 Go de RAM, l'objectif n'est donc pas de lancer toutes les charges en concurrence maximale, mais de contrôler l'activation des modèles, les files, les verrous et les reprises."),
        [
            "Cette contrainte se formalise par un budget mémoire simple. Le système doit conserver une marge pour Windows, l'IDE, le navigateur, Docker Desktop et les services de fond. Dans ce contexte, le raisonnement d'architecture ne consiste pas à maximiser le parallélisme, mais à maintenir la charge utile sous un plafond observable.",
        ],
    )
    eq_anchor = find_contains(doc, "Cette contrainte se formalise par un budget mémoire simple")
    eq = insert_equation_after(eq_anchor, "M_total = M_base + M_Paddle + M_GLM = 2,5 Go + 2,5 Go + 3,5 Go = 8,5 Go")
    fig_anchor = insert_figure_after(
        eq,
        FIG_TECH,
        "Figure 6. Budget mémoire, logique de seuils et amélioration OCR",
        "Source : auteur, d'après docs/analysis-issue-9-ports-wsl2.md, code/backend/app/core/config.py et docs/test-evidence/ocr-beta-loop/report.md.",
        6.2,
    )
    insert_paragraphs_after(
        fig_anchor,
        [
            "La figure 6 donne une lecture synthétique du raisonnement. Le côté gauche rappelle que la charge IA ne doit pas saturer la machine de développement. Le côté droit montre que l'amélioration OCR résulte d'une boucle d'observation, correction et vérification, et non d'un simple appel à une bibliothèque.",
        ],
    )

    replace_exact(
        doc,
        "Le recours à GLM OCR est pertinent pour les cas où la structure du document ou la qualité de l'image rendent la lecture classique insuffisante. Ce type de moteur peut mieux exploiter le contexte visuel et la disposition du document. En contrepartie, il est plus coûteux et doit donc être utilisé avec discernement, par exemple sur les échecs PaddleOCR, les documents complémentaires ou les captures nécessitant une interprétation plus riche.",
        "Le recours à GLM OCR est pertinent pour les cas où la structure du document ou la qualité de l'image rendent la lecture classique insuffisante. Le rapport technique de GLM OCR décrit un modèle multimodal compact, fondé sur un pipeline à deux étages et sur de la reconnaissance régionale parallèle, avec une logique de multi token prediction destinée à améliorer le débit de décodage (Duan et al., 2026 ; zai org, 2026). Dans VeriPass, cette famille de moteur est donc traitée comme un renfort conditionnel, plus coûteux qu'une première passe PaddleOCR, et non comme un remplacement systématique de l'OCR local.",
    )

    insert_paragraphs_after(
        find_exact(doc, "Chaque champ extrait est associé à un score de confiance. Lorsque le score est suffisant, la valeur peut être présentée au client pour confirmation. Lorsque le score est faible, la solution doit conserver l'incertitude : statut PARTIAL ou FAILED, message de reprise, traitement différé ou revue manuelle. Cette gestion explicite des échecs est plus importante qu'un affichage artificiellement confiant."),
        [
            "La granularité des seuils traduit ce raisonnement. Le seuil OCR d'acceptation est fixé à 0,90 dans la configuration, tandis que le seuil d'édition utilisateur est fixé à 0,95. Le premier seuil sert à déclencher un renfort ou une revue, le second à éviter qu'un champ encore fragile soit verrouillé trop tôt.",
        ],
    )
    score_anchor = find_contains(doc, "La granularité des seuils traduit ce raisonnement")
    score_eq = insert_equation_after(score_anchor, "S_ocr(d) = somme(w_f c_f) / somme(w_f)")
    score_eq = insert_equation_after(score_eq, "GLM(d) = 1 si min(c_f) < 0,90 ou n_champs(d) < 2")
    insert_equation_after(score_eq, "Edition_client(d) = 1 si S_ocr(d) < 0,95")

    insert_paragraphs_after(
        find_exact(doc, "La trajectoire réellement utilisée par le prototype distingue la préparation, la revue et la décision. Le dossier démarre en DRAFT, passe en PENDING_AGENT_REVIEW après soumission, peut revenir en PENDING_INFO lorsqu'un complément est demandé, puis se termine par APPROVED, REJECTED ou FRAUD_SUSPECT selon la décision back office. Cette machine d'états évite de confondre un dossier encore éditable, un dossier en attente de revue et un dossier déjà décidé."),
        [
            "Les états dégradés complètent cette trajectoire. LOCKED_LIVENESS matérialise une suspension temporaire liée aux échecs de vivacité, afin d'éviter les essais répétés et les attaques par présentation. ABANDONED permet de qualifier les parcours interrompus, ce qui donne à la banque un indicateur de friction et évite de mélanger abandon utilisateur, rejet conformité et panne technique.",
        ],
    )

    insert_paragraphs_after(
        find_exact(doc, "La cohérence des états est un point critique. Un dossier ne doit pas être à la fois éditable par le client et approuvé par le back office. Une pièce supprimée ne doit pas rester référencée comme preuve active. Une correction OCR ne doit pas être perdue lorsque le dossier change d'état. Ces règles semblent simples, mais elles exigent une base capable de relier les actions et de rejeter les transitions incohérentes."),
        [
            "Ce choix rejoint le cours de référentiel de données : un objet métier doit être défini de façon stable et partagé par les acteurs. Dans VeriPass, le dossier KYC, le document, le champ OCR, le résultat biométrique et la décision sont des objets métier persistants. Les événements analytiques et indicateurs peuvent ensuite être reconstruits, mais ils ne remplacent pas ces objets de référence.",
        ],
    )

    insert_paragraphs_after(
        find_exact(doc, "Le parallélisme a dû être conçu de manière sélective. Certaines tâches peuvent être traitées en file sans risque, comme les notifications ou des traitements légers. D'autres, notamment GLM OCR et certains contrôles d'image, doivent être limités par des verrous et par une affectation selon la charge active. Le principe retenu rapproche l'affectation de la logique Weighted Round Robin et Least Connections : répartir les dossiers quand cela est possible, mais bloquer la concurrence lorsque la mémoire ou la cohérence du dossier l'exige."),
        [
            "Le routage de charge peut être décrit sans révéler tous les détails d'implémentation. Chaque agent ou worker reçoit une charge calculée à partir du nombre de dossiers en file, du temps moyen de traitement et d'un coefficient de capacité. Le dossier est affecté à la charge normalisée la plus faible, sauf lorsqu'un verrou Redis impose une exécution exclusive.",
        ],
    )
    load_anchor = find_contains(doc, "Le routage de charge peut être décrit")
    load_eq = insert_equation_after(load_anchor, "Charge(a) = q_a / w_a + lambda t_a + rho r_a")
    insert_equation_after(load_eq, "a_choisi = argmin Charge(a)")

    insert_paragraphs_after(
        find_exact(doc, "L'intégration AML/CFT donne au back office une profondeur supplémentaire. Les alertes de sanctions, les conflits de NIU, les documents expirants ou les suspicions de doublon ne relèvent pas du même traitement qu'une simple erreur de saisie. VeriPass prépare donc des signaux distincts afin que chaque rôle puisse traiter le dossier selon sa responsabilité."),
        [
            "Il faut toutefois distinguer trois niveaux. Le prototype implémente la revue de dossier, la décision motivée et certains signaux de risque. Le modèle de données prépare les alertes AML/CFT, les conflits, les listes internes et les contrôles de doublon. La mise en production devra encore valider les connexions aux référentiels officiels, les règles de filtrage et les procédures d'escalade conformité.",
        ],
    )

    insert_paragraphs_after(
        find_exact(doc, "L'évaluation suit une logique de traçabilité. Chaque preuve doit être reliée à une exigence : la soumission vérifie la complétude du parcours, l'OCR vérifie l'extraction documentaire, le liveness vérifie l'intégration biométrique, le back office vérifie la décision humaine, et l'audit log vérifie la reconstitution des actions. Cette lecture évite de présenter des tests isolés sans rapport avec la problématique."),
        [
            "La traçabilité peut être lue comme une chaîne exigence, composant et preuve. La capture CNI produit un fichier, un hash et un document associé. L'OCR produit des champs, des scores et des corrections possibles. Le liveness produit un score, un statut et un motif. La soumission produit une transition d'état. La décision back office produit un acteur, un motif et une entrée d'audit. Ce chaînage rend l'évaluation plus solide qu'une simple démonstration d'écran.",
        ],
    )

    insert_paragraphs_after(
        find_exact(doc, "La biométrie doit donc rester cadrée comme un signal d'aide à la revue. Le score de liveness et le statut de face matching donnent au back office une information supplémentaire, mais ne remplacent pas la décision humaine. Cette position réduit le risque de rejet injustifié et maintient la responsabilité métier au niveau des acteurs habilités."),
        [
            "La conservation biométrique doit rester gouvernée par une politique explicite. Le prototype traite les images et conserve les résultats utiles à l'audit technique ; la conservation durable du selfie ou des preuves brutes devra être décidée par la conformité, selon la finalité, la durée de rétention et les droits des personnes. Cette distinction évite de confondre preuve de faisabilité et politique de production.",
        ],
    )

    insert_paragraphs_after(
        find_exact(doc, "Sur le plan financier, l'impact devra être évalué à partir d'indicateurs de pilote : temps moyen de constitution, taux de reprises, temps agent par dossier, coût de stockage, consommation IA, coût de connexion, énergie et maintenance. Le mémoire ne fixe pas un retour sur investissement artificiel, mais il fournit une méthode pour comparer le coût de reproduction du prototype avec une prestation équivalente du marché sur un périmètre comparable."),
        [
            "Un poste peut déjà être chiffré de façon objective : le transport. En retenant l'hypothèse demandée de 2 200 FCFA par aller retour, cinq jours par semaine, du 19 janvier 2026 au 19 juin 2026 inclus, on obtient 110 jours ouvrés, hors jours fériés non déduits. Le coût transport estimé est donc de 242 000 FCFA.",
        ],
    )
    cost_anchor = find_contains(doc, "Un poste peut déjà être chiffré")
    cost_eq = insert_equation_after(cost_anchor, "C_transport = 110 x 2 200 = 242 000 FCFA")
    cost_eq = insert_equation_after(cost_eq, "C_total_projet = C_transport + C_connexion + C_electricite + C_machine + C_assistance_IA + C_temps_humain + C_hebergement_pilote")
    fig_cost = insert_figure_after(
        cost_eq,
        FIG_COST,
        "Figure 7. Coût transport cumulé du stage selon l'hypothèse cinq jours par semaine",
        "Source : calcul de l'auteur, sur la base d'un aller retour quotidien estimé à 2 200 FCFA.",
        6.2,
    )
    insert_paragraphs_after(
        fig_cost,
        [
            "La figure 7 n'épuise pas l'analyse financière, mais elle montre la méthode attendue : identifier un poste, fixer une hypothèse, expliciter le calcul, puis distinguer les coûts chiffrés des coûts à renseigner. Les autres postes, connexion, électricité, amortissement machine, assistance IA, supervision et temps humain, devront être complétés avec des justificatifs ou des hypothèses validées.",
        ],
    )

    replace_exact(
        doc,
        "La perspective la plus immédiate consiste à transformer les preuves techniques en indicateurs de pilotage. Il faudra mesurer le temps moyen de constitution, le taux de dossiers incomplets, le nombre de corrections OCR, les motifs de demande de complément, les cas de rejet biométrique et la charge par rôle back office. Ces indicateurs permettront de juger la solution sur son utilité réelle, et non sur la seule existence des composants techniques.",
        "La perspective la plus immédiate consiste à transformer les preuves techniques en indicateurs de pilotage. Il faudra mesurer le temps moyen de constitution, le taux de dossiers incomplets, le nombre de corrections OCR, les motifs de demande de complément, les cas de rejet biométrique, la charge par rôle back office et le coût par dossier instruit. Ces indicateurs permettront de juger la solution sur son utilité réelle, et non sur la seule existence des composants techniques.",
    )

    replace_exact(
        doc,
        "Les perspectives d'industrialisation sont néanmoins claires. Elles concernent le durcissement sécurité, la validation conformité, la supervision, la calibration sur données représentatives, l'amélioration de l'expérience mobile, l'extension des contrôles AML/CFT et la mesure des gains opérationnels. À terme, VeriPass peut devenir un socle RegTech pour l'onboarding bancaire dans la zone CEMAC, à condition de conserver son principe directeur : automatiser ce qui peut être structuré, journaliser ce qui doit être prouvé, et laisser aux acteurs habilités la responsabilité des décisions sensibles.",
        "Les perspectives d'industrialisation sont néanmoins claires. Elles concernent le durcissement sécurité, la validation conformité, la supervision, la calibration sur données représentatives, l'amélioration de l'expérience mobile, l'extension des contrôles AML/CFT et la mesure des gains opérationnels. À terme, VeriPass peut constituer une base de travail pour un socle RegTech d'onboarding bancaire dans la zone CEMAC, sous réserve d'un pilote, d'une calibration documentée et d'une validation conformité. Son principe directeur reste le même : automatiser ce qui peut être structuré, journaliser ce qui doit être prouvé, et laisser aux acteurs habilités la responsabilité des décisions sensibles.",
    )

    add_bibliography_entries(doc)
    final_checks(doc)
    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
