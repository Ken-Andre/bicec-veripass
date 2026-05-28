from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.enum.style import WD_STYLE_TYPE


OUT_DIR = Path(__file__).resolve().parent
WORKSPACE = OUT_DIR.parents[1]
DOCX_PATH = OUT_DIR / "memoire-stage-andre-yoann-kenmogne-v1.docx"
FIG_DIR = OUT_DIR / "figures"
FIG_DIR.mkdir(parents=True, exist_ok=True)

UCAC_LOGO = Path(r"C:\Users\yoann\Downloads\ucac-icam-logo.jpg")
BICEC_LOGO = Path(r"C:\Users\yoann\Downloads\bicec-logo.jpg")

NAVY = RGBColor(31, 58, 95)
BICEC_ORANGE = RGBColor(227, 123, 3)
DARK_BROWN = RGBColor(74, 34, 5)
GRAY = RGBColor(90, 90, 90)
LIGHT_GRAY = "F2F4F7"
LIGHT_ORANGE = "FFF4E8"


def set_run_font(run, name="Arial", size=12, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:cs"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def style_font(style, name="Arial", size=12, color=None, bold=None):
    style.font.name = name
    style._element.rPr.rFonts.set(qn("w:ascii"), name)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    style._element.rPr.rFonts.set(qn("w:cs"), name)
    style.font.size = Pt(size)
    if color is not None:
        style.font.color.rgb = color
    if bold is not None:
        style.font.bold = bold


def configure_styles(doc):
    styles = doc.styles

    normal = styles["Normal"]
    style_font(normal, "Arial", 12, RGBColor(0, 0, 0))
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(0)

    body = styles.add_style("Memoire Body", WD_STYLE_TYPE.PARAGRAPH)
    style_font(body, "Arial", 12, RGBColor(0, 0, 0))
    body.paragraph_format.first_line_indent = Inches(0.5)
    body.paragraph_format.line_spacing = 1.5
    body.paragraph_format.space_after = Pt(6)
    body.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    front = styles.add_style("Front Matter", WD_STYLE_TYPE.PARAGRAPH)
    style_font(front, "Arial", 12, RGBColor(0, 0, 0))
    front.paragraph_format.line_spacing = 1.5
    front.paragraph_format.space_after = Pt(6)
    front.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    quote = styles.add_style("Academic Quote", WD_STYLE_TYPE.PARAGRAPH)
    style_font(quote, "Arial", 11, GRAY)
    quote.paragraph_format.left_indent = Inches(0.5)
    quote.paragraph_format.right_indent = Inches(0.5)
    quote.paragraph_format.line_spacing = 1.0
    quote.paragraph_format.space_after = Pt(6)
    quote.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for name, size, color, before, after in [
        ("Title", 18, NAVY, 0, 8),
        ("Heading 1", 16, NAVY, 16, 8),
        ("Heading 2", 14, NAVY, 12, 6),
        ("Heading 3", 12, DARK_BROWN, 8, 4),
    ]:
        style = styles[name]
        style_font(style, "Arial", size, color, True)
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    caption = styles.add_style("Caption Memoire", WD_STYLE_TYPE.PARAGRAPH)
    style_font(caption, "Arial", 11, GRAY)
    caption.paragraph_format.line_spacing = 1.0
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    source = styles.add_style("Source Memoire", WD_STYLE_TYPE.PARAGRAPH)
    style_font(source, "Arial", 10, GRAY)
    source.paragraph_format.line_spacing = 1.0
    source.paragraph_format.space_before = Pt(2)
    source.paragraph_format.space_after = Pt(8)

    bibliography = styles.add_style("Bibliographie", WD_STYLE_TYPE.PARAGRAPH)
    style_font(bibliography, "Arial", 10, RGBColor(0, 0, 0))
    bibliography.paragraph_format.line_spacing = 1.0
    bibliography.paragraph_format.space_after = Pt(6)
    bibliography.paragraph_format.hanging_indent = Inches(0.35)


def set_section_geometry(section):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)


def add_field(paragraph, instr):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = instr
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr_text)
    run._r.append(fld_end)


def set_page_numbering(section, fmt="decimal", start=1, show=True):
    section.footer.is_linked_to_previous = False
    sect_pr = section._sectPr
    existing = sect_pr.find(qn("w:pgNumType"))
    if existing is not None:
        sect_pr.remove(existing)
    pg_num = OxmlElement("w:pgNumType")
    pg_num.set(qn("w:start"), str(start))
    pg_num.set(qn("w:fmt"), fmt)
    sect_pr.append(pg_num)

    footer = section.footer
    for paragraph in footer.paragraphs:
        paragraph.clear()
    if show:
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_field(p, "PAGE")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge in kwargs:
            tag = "w:{}".format(edge)
            element = tc_borders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                tc_borders.append(element)
            for key, value in kwargs[edge].items():
                element.set(qn("w:{}".format(key)), str(value))


def remove_table_borders(table):
    for row in table.rows:
        for cell in row.cells:
            set_cell_border(
                cell,
                top={"val": "nil"},
                bottom={"val": "nil"},
                left={"val": "nil"},
                right={"val": "nil"},
                insideH={"val": "nil"},
                insideV={"val": "nil"},
            )


def keep_table_rows_together(table):
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        if tr_pr.find(qn("w:cantSplit")) is None:
            tr_pr.append(OxmlElement("w:cantSplit"))
    if table.rows:
        tr_pr = table.rows[0]._tr.get_or_add_trPr()
        if tr_pr.find(qn("w:tblHeader")) is None:
            tr_pr.append(OxmlElement("w:tblHeader"))


def format_table(table, header_fill=LIGHT_GRAY, font_size=10):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    keep_table_rows_together(table)
    for row_idx, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_after = Pt(3)
                for run in p.runs:
                    set_run_font(run, size=font_size)
            if row_idx == 0:
                set_cell_shading(cell, header_fill)
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.bold = True


def add_cover(doc):
    section = doc.sections[0]
    set_section_geometry(section)
    set_page_numbering(section, show=False)

    top = doc.add_table(rows=1, cols=2)
    remove_table_borders(top)
    top.columns[0].width = Inches(2.2)
    top.columns[1].width = Inches(2.2)
    left_p = top.cell(0, 0).paragraphs[0]
    right_p = top.cell(0, 1).paragraphs[0]
    left_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    right_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if BICEC_LOGO.exists():
        left_p.add_run().add_picture(str(BICEC_LOGO), width=Inches(0.95))
    else:
        left_p.add_run("BICEC")
    if UCAC_LOGO.exists():
        right_p.add_run().add_picture(str(UCAC_LOGO), width=Inches(1.55))
    else:
        right_p.add_run("UCAC-ICAM")

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("MEMOIRE DE FIN D'ETUDES D'INGENIEUR")
    set_run_font(r, size=17, color=NAVY, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Cycle ingénieur X5 - UCAC-ICAM")
    set_run_font(r, size=13, color=GRAY, bold=True)

    doc.add_paragraph()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(18)
    r = title.add_run(
        "Conception et développement d'un écosystème intelligent d'acquisition client :\n"
        "intégration de mécanismes avancés de vérification pour l'automatisation\n"
        "de la conformité Know Your Customer (KYC)"
    )
    set_run_font(r, size=19, color=DARK_BROWN, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = subtitle.add_run("Projet réalisé au sein de la BICEC")
    set_run_font(r, size=13, color=NAVY, bold=True)

    meta = doc.add_table(rows=6, cols=2)
    remove_table_borders(meta)
    rows = [
        ("Etudiant", "Andre Yoann KENMOGNE"),
        ("Entreprise d'accueil", "Banque Internationale du Cameroun pour l'Epargne et le Crédit (BICEC)"),
        ("Département", "Etude & Développement"),
        ("Tuteur entreprise", "M. Jackson Parfait KOMBE LELE"),
        ("Encadreur institut", "[Nom de l'encadreur UCAC-ICAM à compléter]"),
        ("Période de stage", "du 19 janvier 2026 au 18 juin 2026"),
    ]
    for idx, (label, value) in enumerate(rows):
        meta.cell(idx, 0).text = label
        meta.cell(idx, 1).text = value
    for row in meta.rows:
        row.cells[0].width = Inches(1.8)
        row.cells[1].width = Inches(4.5)
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                for run in p.runs:
                    set_run_font(run, size=11, color=RGBColor(0, 0, 0))
        for run in row.cells[0].paragraphs[0].runs:
            run.bold = True

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Année académique 2025-2026")
    set_run_font(r, size=12, color=GRAY, bold=True)


def add_page_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run(text.upper())
    set_run_font(r, size=16, color=NAVY, bold=True)


def add_body_paragraph(doc, text):
    return doc.add_paragraph(text, style="Memoire Body")


def add_front_paragraph(doc, text):
    return doc.add_paragraph(text, style="Front Matter")


def add_caption_above(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_run_font(r, size=11, color=RGBColor(0, 0, 0), bold=True)
    return p


def add_figure_caption(doc, text):
    p = doc.add_paragraph(style="Caption Memoire")
    r = p.add_run(text)
    set_run_font(r, size=11, color=GRAY, bold=False)
    return p


def add_list(doc, items, numbered=False):
    style = "List Number" if numbered else "List Bullet"
    for item in items:
        p = doc.add_paragraph(item, style=style)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(2)
        for run in p.runs:
            set_run_font(run, size=10.5)


def create_pipeline_figure(path):
    width, height = 1500, 480
    img = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)
    try:
        font_title = ImageFont.truetype(r"C:\Windows\Fonts\arialbd.ttf", 28)
        font = ImageFont.truetype(r"C:\Windows\Fonts\arial.ttf", 22)
        font_small = ImageFont.truetype(r"C:\Windows\Fonts\arial.ttf", 18)
    except OSError:
        font_title = font = font_small = None

    draw.text((40, 25), "Chaîne fonctionnelle du pipeline KYC VeriPass", fill=(31, 58, 95), font=font_title)
    boxes = [
        ("Capture\nmobile", "CNI, selfie,\nNIU, facture"),
        ("Prétraitement\n& OCR", "PaddleOCR,\nGLM-OCR"),
        ("Contrôles\nbiométriques", "Liveness,\nface match"),
        ("Stockage\n& traces", "PostgreSQL,\nSHA-256, audit"),
        ("Scoring\n& routage", "Readiness,\npriorité revue"),
        ("Validation\nhumaine", "Jean, Thomas,\nSylvie"),
    ]
    x0, y0, bw, bh, gap = 45, 120, 205, 170, 35
    for i, (title, sub) in enumerate(boxes):
        x = x0 + i * (bw + gap)
        fill = (255, 244, 232) if i % 2 == 0 else (242, 244, 247)
        outline = (227, 123, 3) if i % 2 == 0 else (31, 58, 95)
        draw.rounded_rectangle([x, y0, x + bw, y0 + bh], radius=18, fill=fill, outline=outline, width=3)
        draw.text((x + 18, y0 + 24), title, fill=(74, 34, 5), font=font_title, spacing=4)
        draw.text((x + 18, y0 + 95), sub, fill=(60, 60, 60), font=font_small, spacing=4)
        if i < len(boxes) - 1:
            ax1 = x + bw + 5
            ay = y0 + bh // 2
            ax2 = x + bw + gap - 8
            draw.line([ax1, ay, ax2, ay], fill=(31, 58, 95), width=4)
            draw.polygon([(ax2, ay), (ax2 - 14, ay - 9), (ax2 - 14, ay + 9)], fill=(31, 58, 95))
    draw.text(
        (45, 330),
        "Principe directeur : l'automatisation prépare le dossier et signale les risques ; la décision d'activation reste humaine et auditée.",
        fill=(31, 58, 95),
        font=font,
    )
    img.save(path)


def add_preliminaries(doc):
    prelim = doc.add_section(WD_SECTION.NEW_PAGE)
    set_section_geometry(prelim)
    set_page_numbering(prelim, fmt="lowerRoman", start=1, show=True)

    add_page_title(doc, "Dédicace")
    p = doc.add_paragraph(style="Front Matter")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(
        "Je dédie ce travail à ma famille, pour son soutien constant, sa confiance et les sacrifices "
        "qui ont accompagné mon parcours d'ingénieur."
    )
    set_run_font(r, size=13)
    p = doc.add_paragraph(style="Front Matter")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(
        "Je le dédie également à toutes les personnes qui m'ont appris que la rigueur, la persévérance "
        "et le sens du service sont les fondations d'un projet utile."
    )
    set_run_font(r, size=13)

    doc.add_page_break()
    add_page_title(doc, "Remerciements")
    for text in [
        "Je tiens à exprimer ma reconnaissance à la Banque Internationale du Cameroun pour l'Epargne et le Crédit (BICEC) pour m'avoir accueilli dans un environnement professionnel exigeant, au contact direct d'un projet bancaire à forte portée opérationnelle et réglementaire.",
        "Mes remerciements vont particulièrement à M. Jackson Parfait KOMBE LELE, Responsable du Département Etude & Développement, pour son encadrement, sa disponibilité et la confiance accordée dans la réalisation de ce projet. Je remercie également les équipes de la BICEC qui ont contribué, par leurs échanges et leurs exigences métier, à donner au projet une orientation concrète et utile.",
        "J'adresse aussi mes remerciements à l'UCAC-ICAM et à mes encadreurs académiques pour la formation reçue, l'accompagnement méthodologique et l'exigence qui structure ce mémoire de fin d'études.",
        "Enfin, je remercie ma famille, mes proches et toutes les personnes qui m'ont soutenu durant cette période de stage. Leur présence a été déterminante dans la conduite de ce travail."
    ]:
        add_front_paragraph(doc, text)

    doc.add_page_break()
    add_page_title(doc, "Résumé")
    for text in [
        "Ce mémoire présente la conception et le développement de BICEC VeriPass, un écosystème intelligent d'acquisition client destiné à digitaliser l'entrée en relation bancaire et à renforcer l'automatisation contrôlée de la conformité Know Your Customer (KYC). Le projet s'inscrit dans un contexte de transformation du secteur bancaire camerounais, marqué par la pression concurrentielle des fintechs et du Mobile Money, par les attentes de rapidité des clients et par le renforcement des exigences de vigilance imposées aux établissements assujettis.",
        "La solution proposée vise à réduire les délais de constitution du dossier client, à limiter les erreurs de saisie, à améliorer la traçabilité des contrôles et à fournir aux agents de conformité un dossier numérique complet. Elle combine une application mobile de capture, un backend FastAPI, une base PostgreSQL, des traitements asynchrones Redis/Celery, des mécanismes OCR et biométriques, ainsi qu'un backoffice de validation humaine. L'automatisation n'a pas pour objectif de supprimer le contrôle humain ; elle prépare, structure et qualifie le dossier afin que la décision finale reste explicable, documentée et auditée.",
        "Le mémoire analyse d'abord le contexte institutionnel, réglementaire et opérationnel du projet, puis présente l'architecture fonctionnelle et technique de la solution. Il met en évidence la façon dont les choix d'implémentation répondent aux enjeux de souveraineté des données, d'auditabilité, de sécurité documentaire et d'expérience client."
    ]:
        add_front_paragraph(doc, text)
    p = doc.add_paragraph(style="Front Matter")
    r = p.add_run("Mots-clés : ")
    set_run_font(r, size=12, bold=True)
    p.add_run("KYC, COBAC, onboarding digital, OCR, biométrie, conformité bancaire, FastAPI, React PWA, audit trail, BICEC.")

    doc.add_page_break()
    add_page_title(doc, "Abstract")
    for text in [
        "This report presents the design and development of BICEC VeriPass, an intelligent customer acquisition ecosystem built to digitize bank onboarding and strengthen controlled automation of Know Your Customer (KYC) compliance. The project takes place in the Cameroonian banking context, where banks face stronger competition from fintech and mobile money services, increasing customer expectations for speed, and stricter due diligence obligations.",
        "The proposed solution reduces dossier preparation time, limits manual data entry errors, improves compliance traceability, and provides backoffice agents with structured digital evidence. It combines a mobile capture journey, a FastAPI backend, PostgreSQL persistence, Redis/Celery asynchronous processing, OCR and biometric verification mechanisms, and a human validation backoffice. Automation does not replace the compliance decision; it prepares and qualifies the dossier so that final activation remains explainable, documented, and auditable.",
        "The report first analyzes the institutional, regulatory and operational context, then presents the functional and technical architecture of the solution. It highlights how the implementation choices answer data sovereignty, auditability, document security and customer experience requirements."
    ]:
        add_front_paragraph(doc, text)
    p = doc.add_paragraph(style="Front Matter")
    r = p.add_run("Keywords: ")
    set_run_font(r, size=12, bold=True)
    p.add_run("KYC, COBAC, digital onboarding, OCR, biometrics, banking compliance, FastAPI, React PWA, audit trail, BICEC.")

    doc.add_page_break()
    add_page_title(doc, "Sommaire provisoire")
    add_front_paragraph(
        doc,
        "Le sommaire ci-dessous reprend la structure complète du mémoire. Les numéros de page seront actualisés automatiquement dans la version finale, après stabilisation de toutes les sections et annexes."
    )
    add_caption_above(doc, "Tableau 1. Structure prévue du mémoire")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Partie"
    table.rows[0].cells[1].text = "Contenu"
    rows = [
        ("Introduction générale", "Contexte, problématique, objectifs, méthodologie et plan du mémoire"),
        ("Chapitre 1", "Réglementation COBAC, présentation de la BICEC et analyse du processus KYC existant"),
        ("Chapitre 2", "Analyse fonctionnelle, architecture cible et conception du pipeline de données KYC"),
        ("Chapitre 3", "Implémentation, tests, résultats et limites de la solution BICEC VeriPass"),
        ("Conclusion générale", "Bilan du stage, apports, limites et perspectives"),
        ("Annexes", "Sujet validé, extraits d'architecture, captures, contrats API et preuves de tests"),
    ]
    for left, right in rows:
        cells = table.add_row().cells
        cells[0].text = left
        cells[1].text = right
    format_table(table)

    doc.add_page_break()
    add_page_title(doc, "Listes préliminaires")
    doc.add_heading("Liste des tableaux", level=2)
    for item in [
        "Tableau 1. Structure prévue du mémoire",
        "Tableau 2. Abréviations utilisées",
        "Tableau 3. Fiche signalétique synthétique de la BICEC",
        "Tableau 4. Traduction des exigences réglementaires en mécanismes projet",
        "Tableau 5. Limites du processus manuel et réponse attendue du pipeline numérique",
    ]:
        doc.add_paragraph(item, style="Front Matter")
    doc.add_heading("Liste des figures", level=2)
    doc.add_paragraph("Figure 1. Chaîne fonctionnelle du pipeline KYC VeriPass", style="Front Matter")

    doc.add_page_break()
    add_page_title(doc, "Liste des abréviations")
    add_caption_above(doc, "Tableau 2. Abréviations utilisées")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Abréviation"
    table.rows[0].cells[1].text = "Signification"
    rows = [
        ("AML/CFT", "Anti-Money Laundering / Countering the Financing of Terrorism"),
        ("API", "Application Programming Interface"),
        ("BEAC", "Banque des Etats de l'Afrique Centrale"),
        ("BICEC", "Banque Internationale du Cameroun pour l'Epargne et le Crédit"),
        ("CEMAC", "Communauté Economique et Monétaire de l'Afrique Centrale"),
        ("CNI", "Carte Nationale d'Identité"),
        ("COBAC", "Commission Bancaire de l'Afrique Centrale"),
        ("KYC", "Know Your Customer"),
        ("LBC/FT", "Lutte contre le blanchiment de capitaux et le financement du terrorisme"),
        ("NIU", "Numéro d'Identifiant Unique"),
        ("OCR", "Optical Character Recognition"),
        ("PFE", "Projet de Fin d'Etudes"),
        ("PWA", "Progressive Web Application"),
        ("RBAC", "Role-Based Access Control"),
    ]
    for abbr, meaning in rows:
        cells = table.add_row().cells
        cells[0].text = abbr
        cells[1].text = meaning
    format_table(table)

    doc.add_page_break()
    add_page_title(doc, "Glossaire")
    gloss = [
        ("Audit trail", "Journal horodaté des actions réalisées sur un dossier, permettant de reconstituer qui a fait quoi, quand et pourquoi."),
        ("Backoffice", "Interface interne utilisée par les agents BICEC pour traiter, contrôler et décider sur les dossiers clients."),
        ("Face matching", "Comparaison biométrique entre le visage capturé lors du selfie et la photographie du document d'identité."),
        ("Liveness", "Vérification de vivacité destinée à s'assurer que la personne filmée est présente physiquement et qu'il ne s'agit pas d'une photo, d'un écran ou d'une attaque de présentation."),
        ("Pipeline KYC", "Chaîne de traitement qui collecte les pièces, extrait les données, contrôle leur cohérence, qualifie le risque, stocke les preuves et route le dossier vers la validation humaine."),
        ("Souveraineté des données", "Principe selon lequel les données sensibles restent maîtrisées localement, avec des choix d'hébergement, de stockage et de traitement compatibles avec les contraintes de l'organisation et du pays."),
    ]
    for term, definition in gloss:
        p = doc.add_paragraph(style="Front Matter")
        r = p.add_run(term + " : ")
        set_run_font(r, bold=True)
        p.add_run(definition)


def add_body(doc):
    body_section = doc.add_section(WD_SECTION.NEW_PAGE)
    set_section_geometry(body_section)
    set_page_numbering(body_section, fmt="decimal", start=1, show=True)

    doc.add_heading("Introduction générale", level=1)
    intro_paragraphs = [
        "Le secteur bancaire camerounais est confronté à une double transformation. D'un côté, les clients, en particulier les jeunes actifs, les entrepreneurs et les utilisateurs déjà familiers du Mobile Money, attendent des parcours simples, rapides et disponibles en dehors des horaires d'agence. De l'autre, les banques doivent renforcer leurs dispositifs de conformité, de connaissance client et de traçabilité afin de répondre aux exigences de la Commission Bancaire de l'Afrique Centrale et aux standards de lutte contre le blanchiment de capitaux et le financement du terrorisme.",
        "Dans ce contexte, l'entrée en relation bancaire devient un moment critique. Elle doit être fluide pour le client, mais suffisamment robuste pour l'établissement. Le processus traditionnel, fondé sur la présence en agence, la manipulation de formulaires papier, la photocopie de pièces et la saisie manuelle, crée des délais, des erreurs et une difficulté de pilotage. Les documents peuvent être incomplets, les données saisies peuvent diverger des pièces originales et la reconstitution d'un dossier complet pour un contrôle interne ou externe demande un effort administratif important.",
        "La BICEC, acteur majeur du paysage bancaire camerounais, a donc un intérêt stratégique à digitaliser ce parcours sans fragiliser le contrôle réglementaire. Le sujet de stage validé porte sur la conception et le développement d'un écosystème intelligent d'acquisition client, intégrant des mécanismes avancés de vérification pour automatiser la conformité KYC. Il ne s'agit pas uniquement de créer une application mobile ; il s'agit de construire une chaîne complète qui transforme une demande client en dossier numérique exploitable, vérifiable, stocké, auditable et routé vers les acteurs internes compétents.",
        "Le projet BICEC VeriPass répond à cette ambition. Il propose une Progressive Web App mobile pour le client, une API backend fondée sur FastAPI, une base de données PostgreSQL, des traitements asynchrones Redis/Celery, des mécanismes OCR et biométriques, ainsi qu'un backoffice de validation. Le principe directeur est celui d'une automatisation contrôlée : la machine accélère la capture, l'extraction, la cohérence et la priorisation du dossier, mais la décision d'activation reste humaine. Ce choix est essentiel, car un dossier KYC bancaire ne peut pas être réduit à un score technique ; il doit pouvoir être expliqué, documenté et justifié.",
        "La problématique centrale du mémoire peut donc être formulée ainsi : comment concevoir et implémenter un écosystème numérique d'acquisition client capable de réduire les délais d'onboarding, d'améliorer l'expérience utilisateur et de structurer le contrôle documentaire, tout en maintenant un niveau de conformité, de sécurité et d'auditabilité compatible avec les exigences bancaires et réglementaires de la zone CEMAC ?",
        "Cette problématique se décline en plusieurs questions opérationnelles. Comment collecter les pièces KYC de façon guidée depuis un smartphone ? Comment extraire automatiquement les informations utiles sans supprimer la possibilité de correction humaine ? Comment vérifier la présence du client et la cohérence biométrique ? Comment stocker les preuves et les métadonnées en garantissant l'intégrité du dossier ? Comment organiser une file de validation adaptée aux rôles internes de la banque ? Enfin, comment fournir des indicateurs permettant au management de suivre les délais, les erreurs, les rejets et la charge opérationnelle ?",
        "L'objectif général du stage est de concevoir et développer un prototype avancé de plateforme d'onboarding KYC pour la BICEC. Les objectifs spécifiques sont d'analyser les contraintes réglementaires et métier, de définir un pipeline de données KYC, de mettre en place les interfaces client et backoffice, d'intégrer les briques de vérification documentaire et biométrique, de sécuriser les données sensibles et de produire des traces exploitables pour l'audit.",
        "La méthodologie adoptée combine une analyse documentaire, un cadrage fonctionnel, une conception d'architecture, une implémentation itérative et des tests orientés démonstration métier. Le dépôt projet contient les sources backend, mobile et backoffice, les scripts de déploiement Docker, les contrats API, les modèles de données, les notes de durcissement et les preuves de tests. Ces éléments constituent la base technique exploitée pour rédiger ce mémoire.",
        "Le mémoire est structuré en trois grands chapitres. Le premier présente le cadre d'accueil, le contexte réglementaire COBAC et la problématique opérationnelle du KYC. Le deuxième sera consacré à l'analyse fonctionnelle, à la conception et à l'architecture du pipeline BICEC VeriPass. Le troisième présentera l'implémentation, les tests, les résultats obtenus, les limites observées et les perspectives d'industrialisation. La conclusion reviendra sur les acquis techniques, organisationnels et humains du stage."
    ]
    for text in intro_paragraphs:
        add_body_paragraph(doc, text)

    doc.add_heading("Chapitre 1 - Réglementation COBAC et problématique KYC", level=1)
    add_body_paragraph(
        doc,
        "Ce chapitre établit le cadre dans lequel le projet a été réalisé. Il présente d'abord la BICEC et le contexte du stage, puis analyse les exigences réglementaires liées à la connaissance client. Il met ensuite en évidence les limites du processus manuel d'entrée en relation et formule la problématique à laquelle répond BICEC VeriPass."
    )

    doc.add_heading("1.1 Présentation synthétique de la BICEC et du stage", level=2)
    for text in [
        "La Banque Internationale du Cameroun pour l'Epargne et le Crédit, plus connue sous le sigle BICEC, est une banque camerounaise opérant sur les marchés des particuliers, des professionnels et des entreprises. D'après sa communication institutionnelle, elle est filiale du Groupe Banque Centrale Populaire et figure parmi les acteurs de référence du secteur bancaire au Cameroun. La banque revendique un réseau national de 40 agences, deux centres d'affaires, un espace PME, plus de 380 000 clients et près de 800 collaborateurs (BICEC, 2026).",
        "Cette présence nationale donne à la BICEC une double responsabilité. Elle doit continuer à proposer des services accessibles et proches des clients, tout en modernisant son expérience digitale. Ses valeurs institutionnelles - proximité, citoyenneté, performance et innovation - donnent un cadre favorable à un projet comme VeriPass, dont l'objectif est de transformer l'entrée en relation sans dégrader la confiance bancaire.",
        "Le stage s'est déroulé au sein du Département Etude & Développement, sous l'encadrement de M. Jackson Parfait KOMBE LELE. Le sujet validé par l'entreprise porte sur la conception et le développement d'un écosystème intelligent d'acquisition client, avec intégration de mécanismes avancés de vérification pour l'automatisation de la conformité KYC. La période de stage indiquée sur la fiche de validation s'étend du 19 janvier 2026 au 18 juin 2026."
    ]:
        add_body_paragraph(doc, text)

    add_caption_above(doc, "Tableau 3. Fiche signalétique synthétique de la BICEC")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Elément"
    table.rows[0].cells[1].text = "Information"
    rows = [
        ("Dénomination", "Banque Internationale du Cameroun pour l'Epargne et le Crédit (BICEC)"),
        ("Siège social", "Avenue du Général de Gaulle, Bonanjo, Douala"),
        ("Boîte postale", "B.P. 1925 Douala"),
        ("Téléphone", "(+237) 233 43 60 00"),
        ("Réseau", "40 agences, 2 centres d'affaires et un espace PME"),
        ("Clients et collaborateurs", "Plus de 380 000 clients et près de 800 collaborateurs"),
        ("Tuteur entreprise", "M. Jackson Parfait KOMBE LELE"),
        ("Projet de stage", "Digitalisation de l'onboarding client et automatisation contrôlée du KYC"),
    ]
    for left, right in rows:
        cells = table.add_row().cells
        cells[0].text = left
        cells[1].text = right
    format_table(table, header_fill=LIGHT_ORANGE)
    doc.add_paragraph("Sources : fiche de validation du sujet de stage ; site institutionnel BICEC, pages A propos et Réseau.", style="Source Memoire")

    doc.add_heading("1.2 Concurrence fintech, évolution des usages et pression sur l'entrée en relation", level=2)
    for text in [
        "L'ouverture de compte est un point de contact décisif entre une banque et un futur client. Dans un marché où les services digitaux deviennent familiers, l'utilisateur compare implicitement son expérience bancaire avec les usages qu'il connaît déjà : inscription à un service mobile, création d'un portefeuille électronique, réception instantanée d'un code OTP ou suivi en temps réel d'une demande. Lorsqu'un parcours bancaire exige plusieurs déplacements, des formulaires papier et des délais peu visibles, il devient moins compétitif face aux acteurs digitaux.",
        "Les documents de cadrage du projet identifient un délai d'ouverture pouvant aller de 48 heures à 14 jours dans le processus manuel. Ce délai dépend notamment de la disponibilité du client, de la complétude du dossier, de la qualité des pièces fournies, des allers-retours entre agence et backoffice et de la capacité des agents à traiter les files en attente. Ce modèle crée une friction forte pour les entrepreneurs, les jeunes actifs et les utilisateurs qui ne peuvent pas interrompre facilement leur activité pour se rendre en agence.",
        "La concurrence des fintechs et du Mobile Money ne porte pas seulement sur le prix ou sur la fonctionnalité. Elle porte aussi sur la perception de simplicité. Un service qui permet de commencer une relation à distance, de suivre son statut et de recevoir une réponse rapide crée un avantage de confiance. Pour une banque comme la BICEC, l'enjeu est donc de conserver la rigueur bancaire tout en éliminant les lenteurs qui ne produisent pas de valeur."
    ]:
        add_body_paragraph(doc, text)

    doc.add_heading("1.3 Cadre réglementaire : COBAC R-2023/01 et diligences KYC", level=2)
    for text in [
        "Le règlement COBAC R-2023/01 du 19 décembre 2023 relatif aux diligences des établissements assujettis en matière de lutte contre le blanchiment de capitaux, le financement du terrorisme et de la prolifération constitue un repère majeur pour tout projet d'entrée en relation bancaire dans la zone CEMAC. Il s'applique notamment aux établissements de crédit, aux établissements de microfinance, aux établissements de paiement, aux intermédiaires en opérations de banque et aux bureaux de change. Cette portée confirme que le KYC n'est pas un simple processus commercial : il est une obligation de contrôle.",
        "Le règlement s'inscrit dans une logique d'approche par les risques. Les établissements doivent définir des politiques d'acceptation des clients, identifier les profils susceptibles de présenter un risque et démontrer la pertinence de leur dispositif de gestion et d'atténuation des risques. Cette orientation est essentielle pour VeriPass : le système ne doit pas seulement collecter des pièces ; il doit produire un dossier structuré permettant de comprendre les risques, les contrôles réalisés et les raisons d'une décision.",
        "La connaissance client repose sur l'identification et la vérification de l'identité à partir de documents, données ou informations de sources fiables et indépendantes. Lorsqu'un client souhaite ouvrir un compte, les mesures de vigilance doivent être mises en œuvre. Pour une plateforme numérique, cette exigence se traduit par la capture des pièces, l'extraction des champs utiles, la conservation des images originales, la possibilité de correction humaine et la traçabilité de tout changement apporté aux données extraites.",
        "Le même cadre impose aussi de ne pas tenir de comptes anonymes ou sous des noms fictifs, d'établir le profil de risque du client et de renforcer la vigilance dans certains cas, par exemple lorsque l'opération est inhabituelle, lorsque des informations manquent ou lorsqu'une personne politiquement exposée est détectée. Ces exigences justifient la présence d'un backoffice de conformité, de rôles différenciés et de files de traitement qui séparent la validation KYC, l'analyse AML/CFT et le pilotage opérationnel.",
        "Un autre point déterminant concerne la conservation documentaire. Les commentaires disponibles sur le règlement R-2023/01 indiquent que les documents obtenus à l'ouverture du compte, les éléments de vigilance client, les livres, la correspondance commerciale et les résultats d'analyse doivent être conservés pendant dix ans dans les cas prévus. Pour VeriPass, cette contrainte conduit à une architecture qui stocke les fichiers dans un volume documentaire, conserve les métadonnées et les empreintes SHA-256 en base, et évite de réduire le dossier à une simple donnée OCR."
    ]:
        add_body_paragraph(doc, text)

    add_caption_above(doc, "Tableau 4. Traduction des exigences réglementaires en mécanismes projet")
    table = doc.add_table(rows=1, cols=3)
    for idx, header in enumerate(["Exigence KYC", "Interprétation pour le projet", "Réponse BICEC VeriPass"]):
        table.rows[0].cells[idx].text = header
    rows = [
        ("Identifier le client", "Collecter une preuve d'identité fiable et exploitable", "Capture CNI recto/verso, OCR, conservation du fichier original"),
        ("Vérifier l'identité", "Comparer les informations extraites avec des preuves et signaux complémentaires", "OCR review, liveness, face match, justificatif d'adresse, NIU"),
        ("Appliquer une approche par les risques", "Qualifier les dossiers et signaler les cas à traiter en priorité", "Score de confiance, flags, priorisation backoffice"),
        ("Conserver les documents", "Garder les preuves et les analyses suffisamment longtemps pour audit", "Volume documentaire, métadonnées PostgreSQL, hash SHA-256, backups"),
        ("Tracer les décisions", "Pouvoir reconstituer les actions humaines et système", "Audit log, décisions de validation, correction OCR horodatée"),
        ("Renforcer la vigilance", "Escalader les alertes AML/CFT et les dossiers incohérents", "Rôle Thomas, alertes AML, conflit NIU, revue manuelle renforcée"),
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
    format_table(table)

    doc.add_heading("1.4 Analyse du processus manuel d'entrée en relation", level=2)
    for text in [
        "Le processus manuel d'entrée en relation présente plusieurs limites structurelles. La première est la dépendance à la présence physique. Le client doit se déplacer, parfois plusieurs fois, pour fournir des pièces, corriger une information ou compléter un dossier. Cette contrainte augmente le coût d'opportunité pour les personnes actives et rend l'ouverture de compte moins attractive.",
        "La deuxième limite concerne la qualité des données. Dans un processus papier, les informations sont souvent recopiées depuis des documents dont la lisibilité varie. Une CNI usée, une facture floue ou une attestation partiellement visible peuvent conduire à des erreurs de saisie. Chaque erreur crée un risque opérationnel, car elle peut affecter l'identification du client, la cohérence du dossier ou les contrôles ultérieurs.",
        "La troisième limite porte sur la traçabilité. Lorsque plusieurs personnes consultent, modifient ou complètent un dossier, il devient difficile de reconstituer précisément l'historique si le processus n'est pas outillé. Or, dans un contexte COBAC, la capacité à produire un dossier complet et à justifier une décision est aussi importante que la décision elle-même.",
        "La quatrième limite est organisationnelle. Les agents de conformité doivent consacrer une part importante de leur temps à vérifier la complétude documentaire et à effectuer des relances. Cette charge réduit le temps disponible pour l'analyse à forte valeur ajoutée : cohérence des preuves, identification des anomalies, décision argumentée et suivi des cas sensibles."
    ]:
        add_body_paragraph(doc, text)

    add_caption_above(doc, "Tableau 5. Limites du processus manuel et réponse attendue du pipeline numérique")
    table = doc.add_table(rows=1, cols=3)
    for idx, header in enumerate(["Limite observée", "Risque associé", "Réponse attendue"]):
        table.rows[0].cells[idx].text = header
    rows = [
        ("Déplacements et horaires d'agence", "Abandon client, lenteur, mauvaise expérience", "Parcours mobile accessible à distance"),
        ("Saisie manuelle", "Erreurs, incohérences, corrections tardives", "OCR avec confirmation et correction contrôlée"),
        ("Pièces incomplètes ou illisibles", "Relances, dossier bloqué, non-conformité", "Guidage de capture, readiness check, demande de complément"),
        ("Archivage papier dispersé", "Difficulté d'audit et de reconstitution", "Stockage numérique, hash, métadonnées, audit trail"),
        ("Validation peu pilotable", "Files invisibles, SLA difficile à mesurer", "Backoffice, statuts, analytics et dashboard opérationnel"),
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
    format_table(table, header_fill=LIGHT_ORANGE, font_size=9)

    doc.add_heading("1.5 Problématique retenue", level=2)
    for text in [
        "Le problème à résoudre n'est donc pas seulement technique. Il se situe à l'intersection de l'expérience client, de la conformité réglementaire, de la maîtrise opérationnelle et de la sécurité des données. Une solution trop orientée expérience risquerait de simplifier excessivement le contrôle KYC. A l'inverse, une solution trop orientée conformité pourrait reproduire la lourdeur du papier dans un écran. L'enjeu est de trouver un équilibre.",
        "La problématique retenue pour ce mémoire est la suivante : comment digitaliser l'onboarding client de la BICEC au moyen d'un pipeline intelligent de vérification KYC, afin de réduire les délais et les erreurs de constitution de dossier, tout en garantissant la validation humaine, la conservation des preuves, la sécurité des données et la traçabilité exigée par le cadre bancaire ?",
        "Cette formulation met en avant quatre exigences. La première est la réduction de délai, car le parcours cible vise une expérience inférieure à quinze minutes côté client. La deuxième est la fiabilité documentaire, car l'OCR et la biométrie doivent aider à structurer les preuves. La troisième est la conformité, car le dossier doit pouvoir être audité. La quatrième est l'organisation humaine, car le système doit soutenir les agents au lieu de les contourner."
    ]:
        add_body_paragraph(doc, text)

    doc.add_heading("1.6 Réponse proposée : un pipeline numérique KYC audit-ready", level=2)
    for text in [
        "BICEC VeriPass est conçu comme un pipeline de données et de preuves. La donnée client ne naît pas dans un formulaire isolé ; elle provient d'une chaîne de capture, d'extraction, de vérification, de confirmation, de stockage et de décision. Chaque étape produit des éléments exploitables : un fichier, un hash, des champs OCR, un score, un statut, une action agent ou une entrée d'audit.",
        "Le parcours client mobile permet de démarrer une session KYC, de capturer les documents, de valider les champs extraits, de réaliser un liveness, de fournir les informations d'adresse, de gérer le NIU selon le périmètre MVP, d'accepter les consentements et de signer numériquement. Le backend orchestre ces actions à travers des API versionnées. Le backoffice fournit ensuite aux rôles Jean, Thomas, Sylvie et Admin IT les vues nécessaires à la validation, au contrôle AML/CFT, au pilotage et à l'administration.",
        "L'architecture actuelle du dépôt repose sur React/Vite pour la PWA mobile et le backoffice, FastAPI pour l'API, PostgreSQL pour les données persistantes, Redis et Celery pour les tâches asynchrones, Nginx pour le routage TLS local, et un stockage documentaire externe pour les fichiers KYC. Les modules backend sont organisés par domaine : auth, kyc, backoffice, admin, aml, analytics, audit, banking, devices, notifications et support. Cette organisation rend la solution lisible pour un contexte de stage, tout en préparant une industrialisation progressive."
    ]:
        add_body_paragraph(doc, text)

    fig_path = FIG_DIR / "pipeline_kyc_veripass.png"
    create_pipeline_figure(fig_path)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(fig_path), width=Inches(6.4))
    add_figure_caption(doc, "Figure 1. Chaîne fonctionnelle du pipeline KYC VeriPass")

    doc.add_heading("1.7 Périmètre et délimitations du MVP", level=2)
    for text in [
        "Le périmètre du MVP met l'accent sur la preuve de valeur KYC. La solution démontre la capacité à constituer un dossier numérique complet, à automatiser l'extraction et certains contrôles, à déclencher une revue humaine et à fournir des éléments d'audit. En revanche, elle ne doit pas être présentée comme une application bancaire complète ni comme une intégration production déjà reliée à toutes les API externes.",
        "Les notes de durcissement les plus récentes du projet indiquent que l'intégration DGI réelle est retirée du périmètre MVP. Le NIU est traité comme une preuve locale selon trois états : absent, déclaratif ou document téléversé. De même, les fonctionnalités bancaires présentes dans l'interface servent de démonstration de parcours et de découverte de services ; les opérations réelles de core banking et les transactions financières complètes relèvent d'une phase ultérieure.",
        "Cette délimitation renforce la crédibilité du mémoire. Elle montre que le projet ne confond pas prototype avancé, MVP de conformité et système bancaire production. Le résultat attendu du stage est une base technique et fonctionnelle capable de démontrer l'onboarding KYC, d'appuyer la discussion métier et de préparer les prochaines étapes d'intégration."
    ]:
        add_body_paragraph(doc, text)

    doc.add_heading("1.8 Synthèse du chapitre", level=2)
    for text in [
        "Ce premier chapitre a montré que BICEC VeriPass répond à un besoin réel : réduire la friction de l'ouverture de compte sans affaiblir les obligations de vigilance. La BICEC évolue dans un environnement où l'expérience digitale devient un facteur concurrentiel, mais où la conformité COBAC impose des exigences fortes de connaissance client, de conservation documentaire, d'approche par les risques et de traçabilité.",
        "La problématique du mémoire consiste donc à concevoir une automatisation utile, mais contrôlée. La technologie doit préparer le dossier, améliorer la qualité de la donnée, guider le client et accélérer la revue, tout en laissant à l'agent la responsabilité de la décision finale. Le chapitre suivant développera l'analyse fonctionnelle et l'architecture de cette réponse."
    ]:
        add_body_paragraph(doc, text)

    doc.add_heading("Chapitre 2 - Analyse, conception et architecture de la solution", level=1)
    add_body_paragraph(
        doc,
        "Ce chapitre sera développé dans la prochaine itération du mémoire. Il présentera les besoins fonctionnels, les cas d'utilisation, les acteurs, l'architecture logique, le modèle de données, les contrats API et les choix techniques retenus pour BICEC VeriPass."
    )

    doc.add_heading("Chapitre 3 - Implémentation, tests et évaluation", level=1)
    add_body_paragraph(
        doc,
        "Ce chapitre sera développé dans la prochaine itération du mémoire. Il présentera les modules réalisés, les écrans majeurs, les mécanismes OCR et biométriques, les tests unitaires et end-to-end, les preuves de démonstration, les limites et les perspectives d'industrialisation."
    )

    doc.add_heading("Conclusion générale", level=1)
    add_body_paragraph(
        doc,
        "La conclusion générale sera finalisée après rédaction des chapitres techniques et analyse des résultats. Elle devra positionner le résultat obtenu par rapport aux objectifs du stage, expliciter les apports pour la BICEC, les acquis d'apprentissage et les axes de progrès."
    )

    doc.add_heading("Références bibliographiques provisoires", level=1)
    refs = [
        "BICEC. (2026). A propos de la BICEC. https://www.bicec.com/la-bicec/",
        "BICEC. (2026). Le réseau de la BICEC. https://www.bicec.com/la-bicec/reseau/",
        "COBAC. (2023). Règlement COBAC R-2023/01 du 19 décembre 2023 relatif aux diligences des établissements assujettis en matière de LBC/FT et de prolifération.",
        "Droit Médias Finance. (2024). CEMAC - Le règlement COBAC R-2023/01 relatif aux nouvelles diligences anti-blanchiment entre en vigueur le 1er juillet 2024. https://www.droitmediasfinance.com/",
        "Base Réglementaire. (2025). Articles du règlement COBAC R-2023/01 : articles 1, 3, 11, 12, 13, 38, 44, 45 et 46. https://basereglementaire.com/",
        "UCAC-ICAM. (s. d.). Protocole de rédaction du rapport de stage, d'alternance, du MFH et du stage humanitaire.",
        "UCAC-ICAM. (s. d.). Evaluation du rapport.",
        "BICEC VeriPass. (2026). Documentation interne du projet : cadrage, PRD, architecture, contrats API, modèles de données, notes de durcissement et preuves de tests.",
    ]
    for ref in refs:
        doc.add_paragraph(ref, style="Bibliographie")

    doc.add_heading("Annexes prévues", level=1)
    add_list(
        doc,
        [
            "Annexe A : Fiche de validation du sujet de stage en entreprise.",
            "Annexe B : Architecture conteneurisée de BICEC VeriPass.",
            "Annexe C : Extraits des contrats API KYC et backoffice.",
            "Annexe D : Captures des écrans mobile et backoffice.",
            "Annexe E : Preuves de tests et scénarios de démonstration MVP.",
        ],
        numbered=False,
    )


def build_document():
    doc = Document()
    configure_styles(doc)
    add_cover(doc)
    add_preliminaries(doc)
    add_body(doc)
    doc.core_properties.author = "Andre Yoann KENMOGNE"
    doc.core_properties.title = "Mémoire de fin d'études - BICEC VeriPass"
    doc.core_properties.subject = "Onboarding digital KYC et conformité bancaire"
    doc.core_properties.keywords = "BICEC, KYC, COBAC, OCR, biométrie, FastAPI, React PWA"
    doc.save(DOCX_PATH)
    return DOCX_PATH


if __name__ == "__main__":
    output = build_document()
    print(output)
