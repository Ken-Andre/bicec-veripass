from __future__ import annotations

import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.table import Table


ROOT = Path(__file__).resolve().parents[2]
SRC = ROOT / "docs" / "rapport-stage" / "memoire-stage-andre-yoann-kenmogne-v7.docx"
OUT = ROOT / "docs" / "rapport-stage" / "memoire-stage-andre-yoann-kenmogne-v8.docx"


def set_text(paragraph: Paragraph, text: str) -> None:
    style = paragraph.style
    for run in paragraph.runs:
        run.text = ""
    run = paragraph.add_run(text)
    paragraph.style = style


def insert_after(paragraph: Paragraph, text: str, style: str | None = None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    p = Paragraph(new_p, paragraph._parent)
    if style:
        p.style = style
    p.add_run(text)
    return p


def remove_block(block: Paragraph | Table) -> None:
    el = block._element
    parent = el.getparent()
    parent.remove(el)


def iter_blocks(document: Document):
    for child in document.element.body.iterchildren():
        if child.tag.endswith("p"):
            yield Paragraph(child, document)
        elif child.tag.endswith("tbl"):
            yield Table(child, document)


def table_header(table: Table) -> str:
    if not table.rows:
        return ""
    return " | ".join(cell.text.strip().replace("\n", " ") for cell in table.rows[0].cells)


def paragraphs_by_text(document: Document, text: str) -> list[Paragraph]:
    return [p for p in document.paragraphs if p.text.strip() == text]


def remove_para_with_following_source(document: Document, starts: tuple[str, ...]) -> None:
    blocks = list(iter_blocks(document))
    remove_next_source = False
    to_remove: list[Paragraph] = []
    for block in blocks:
        if isinstance(block, Paragraph):
            text = block.text.strip()
            if remove_next_source and text.startswith("Source :"):
                to_remove.append(block)
                remove_next_source = False
                continue
            remove_next_source = False
            if text.startswith(starts):
                to_remove.append(block)
                remove_next_source = True
    for p in to_remove:
        remove_block(p)


def remove_tables_by_header(document: Document, headers: tuple[str, ...]) -> None:
    for block in list(iter_blocks(document)):
        if isinstance(block, Table):
            head = table_header(block)
            if head.startswith(headers):
                remove_block(block)


def replace_paragraph(document: Document, old: str, new: str) -> bool:
    for p in document.paragraphs:
        if p.text.strip() == old:
            set_text(p, new)
            return True
    return False


def add_section(anchor: Paragraph, heading: str, paragraphs: list[str]) -> Paragraph:
    cur = insert_after(anchor, heading, "Heading 2")
    for text in paragraphs:
        cur = insert_after(cur, text, "Memoire Body")
    return cur


def replace_block_between(document: Document, start_heading: str, end_heading: str, sections: list[tuple[str, list[str]]]) -> None:
    blocks = list(iter_blocks(document))
    start_idx = next(i for i, b in enumerate(blocks) if isinstance(b, Paragraph) and b.text.strip() == start_heading)
    end_idx = next(i for i, b in enumerate(blocks) if i > start_idx and isinstance(b, Paragraph) and b.text.strip() == end_heading)
    start = blocks[start_idx]
    end = blocks[end_idx]
    for b in blocks[start_idx + 1 : end_idx]:
        remove_block(b)
    cur = start
    for heading, paras in sections:
        cur = add_section(cur, heading, paras)


def replace_after_heading_until(document: Document, heading: str, end_heading: str, paragraphs: list[str]) -> None:
    blocks = list(iter_blocks(document))
    start_idx = next(i for i, b in enumerate(blocks) if isinstance(b, Paragraph) and b.text.strip() == heading)
    end_idx = next(i for i, b in enumerate(blocks) if i > start_idx and isinstance(b, Paragraph) and b.text.strip() == end_heading)
    start = blocks[start_idx]
    for b in blocks[start_idx + 1 : end_idx]:
        remove_block(b)
    cur = start
    for text in paragraphs:
        cur = insert_after(cur, text, "Memoire Body")


def replace_between_paragraphs(document: Document, start_text: str, end_text: str, entries: list[str], style: str) -> None:
    blocks = list(iter_blocks(document))
    start_idx = next(i for i, b in enumerate(blocks) if isinstance(b, Paragraph) and b.text.strip() == start_text)
    end_idx = next(i for i, b in enumerate(blocks) if i > start_idx and isinstance(b, Paragraph) and b.text.strip() == end_text)
    start = blocks[start_idx]
    for b in blocks[start_idx + 1 : end_idx]:
        remove_block(b)
    cur = start
    for entry in entries:
        cur = insert_after(cur, entry, style)


def insert_toc_after(paragraph: Paragraph) -> None:
    p = insert_after(paragraph, "", "Normal")
    run = p.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Table des matières générée par Word"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(placeholder)
    run._r.append(fld_end)


def style_major_headings(document: Document) -> None:
    major = (
        "Introduction générale",
        "Chapitre 1 :",
        "Chapitre 2 :",
        "Chapitre 3 :",
        "Chapitre 4 :",
        "Chapitre 5 :",
        "Conclusion générale",
        "Références bibliographiques",
        "Annexes",
    )
    for p in document.paragraphs:
        t = p.text.strip()
        if t == "Introduction générale" or t.startswith(major[1:]):
            p.paragraph_format.page_break_before = True


def main() -> None:
    shutil.copyfile(SRC, OUT)
    doc = Document(OUT)

    replacements = {
        "Ce mémoire présente la conception et le développement de BICEC VeriPass, un écosystème intelligent d'acquisition client destiné à digitaliser l'entrée en relation bancaire et à renforcer l'automatisation contrôlée de la conformité Know Your Customer (KYC). Le projet s'inscrit dans un contexte de transformation du secteur bancaire camerounais, marqué par la pression concurrentielle des fintechs et du Mobile Money, par les attentes de rapidité des clients et par le renforcement des exigences de vigilance imposées aux établissements assujettis.": "Ce mémoire présente la conception et le développement de BICEC VeriPass, une plateforme d'onboarding KYC destinée à structurer l'entrée en relation bancaire à la BICEC. Le projet s'inscrit dans un contexte camerounais où les parcours financiers mobiles habituent les clients à des services plus rapides, tandis que les banques doivent renforcer l'identification, la vérification, la conservation des preuves et la traçabilité exigées par la réglementation COBAC.",
        "La solution proposée vise à réduire les délais de constitution du dossier client, à limiter les erreurs de saisie, à améliorer la traçabilité des contrôles et à fournir aux agents de conformité un dossier numérique complet. Elle combine une application mobile de capture, un backend FastAPI, une base PostgreSQL, des traitements asynchrones Redis/Celery, des mécanismes OCR et biométriques, ainsi qu'un back office de validation humaine. L'automatisation n'a pas pour objectif de supprimer le contrôle humain ; elle prépare, structure et qualifie le dossier afin que la décision finale reste explicable, documentée et auditée.": "La solution proposée vise à réduire les reprises de saisie, les relances liées aux pièces illisibles et la dispersion des preuves. Elle combine une application mobile de capture, une API FastAPI, une base PostgreSQL, des traitements asynchrones Redis/Celery, un pipeline OCR, une vérification de vivacité, une comparaison faciale et un back office de validation. L'automatisation prépare le dossier, mais la décision finale reste portée par un acteur habilité et journalisée dans le système.",
        "This report presents the design and development of BICEC VeriPass, an intelligent customer acquisition ecosystem built to digitize bank onboarding and strengthen controlled automation of Know Your Customer (KYC) compliance. The project takes place in the Cameroonian banking context, where banks face stronger competition from fintech and mobile money services, increasing customer expectations for speed, and stricter due diligence obligations.": "This report presents the design and development of BICEC VeriPass, a KYC onboarding platform built to structure digital customer acquisition at BICEC. The project takes place in the Cameroonian banking context, where mobile financial services raise customer expectations while banks must maintain due diligence, evidence retention, traceability and human review.",
        "The proposed solution reduces dossier preparation time, limits manual data entry errors, improves compliance traceability, and provides back office agents with structured digital evidence. It combines a mobile capture journey, a FastAPI backend, PostgreSQL persistence, Redis/Celery asynchronous processing, OCR and biometric verification mechanisms, and a human validation back office. Automation does not replace the compliance decision; it prepares and qualifies the dossier so that final activation remains explainable, documented, and auditable.": "The proposed solution reduces manual re-entry, limits follow-up requests caused by unreadable documents, and gives back office agents a structured digital dossier. It combines a mobile capture journey, a FastAPI backend, PostgreSQL persistence, Redis/Celery asynchronous processing, OCR, liveness verification, face matching and a human validation back office. Automation does not replace the compliance decision; it prepares and qualifies the dossier so that the final decision remains explainable, documented and auditable.",
        "La BICEC, acteur majeur du paysage bancaire camerounais, a donc un intérêt stratégique à digitaliser ce parcours sans fragiliser le contrôle réglementaire. Le sujet de stage validé porte sur la conception et le développement d'un écosystème intelligent d'acquisition client, intégrant des mécanismes avancés de vérification pour automatiser la conformité KYC. Il ne s'agit pas uniquement de créer une application mobile ; il s'agit de construire une chaîne complète qui transforme une demande client en dossier numérique exploitable, vérifiable, stocké, auditable et routé vers les acteurs internes compétents.": "La BICEC, acteur majeur du paysage bancaire camerounais, a donc un intérêt stratégique à digitaliser ce parcours sans affaiblir le contrôle réglementaire. Le sujet de stage validé porte sur la conception et le développement d'un écosystème d'acquisition client intégrant des mécanismes de vérification KYC. Dans ce mémoire, cette formulation est traitée comme un problème d'ingénierie : transformer une demande client en dossier numérique exploitable, vérifiable, conservé, auditable et routé vers les acteurs internes compétents.",
        "Le projet BICEC VeriPass répond à cette ambition. Il propose une Progressive Web App mobile pour le client, une API backend fondée sur FastAPI, une base de données PostgreSQL, des traitements asynchrones Redis/Celery, des mécanismes OCR et biométriques, ainsi qu'un back office de validation. Le principe directeur est celui d'une automatisation contrôlée : la machine accélère la capture, l'extraction, la cohérence et la priorisation du dossier, mais la décision d'activation reste humaine. Ce choix est essentiel, car un dossier KYC bancaire ne peut pas être réduit à un score technique ; il doit pouvoir être expliqué, documenté et justifié.": "BICEC VeriPass répond à cette ambition par une chaîne composée d'une PWA mobile, d'une API FastAPI, d'une base PostgreSQL, de workers Redis/Celery, d'un stockage documentaire et d'un back office de validation. Le système automatise la capture, l'extraction OCR, certains contrôles de cohérence et la préparation du dossier. La décision d'activation reste humaine, car un dossier KYC bancaire doit pouvoir être expliqué, documenté et justifié.",
        "La problématique centrale du mémoire peut donc être formulée ainsi : comment concevoir et implémenter un écosystème numérique d'acquisition client capable de réduire les délais d'onboarding, d'améliorer l'expérience utilisateur et de structurer le contrôle documentaire, tout en maintenant un niveau de conformité, de sécurité et d'auditabilité compatible avec les exigences bancaires et réglementaires de la zone CEMAC ?": "La problématique centrale du mémoire peut donc être formulée ainsi : comment concevoir et implémenter une plateforme numérique d'onboarding capable de réduire les reprises et les délais de constitution d'un dossier KYC, tout en maintenant une conformité, une sécurité et une auditabilité compatibles avec les exigences bancaires de la zone CEMAC ?",
        "Le mémoire analyse d'abord le contexte institutionnel, réglementaire et opérationnel du projet, puis présente l'architecture fonctionnelle et technique de la solution. Il met en évidence la façon dont les choix d'implémentation répondent aux enjeux de souveraineté des données, d'auditabilité, de sécurité documentaire et d'expérience client.": "Le mémoire analyse d'abord le contexte institutionnel, réglementaire et opérationnel du projet, puis présente les besoins, les données et l'architecture de traitement. Il décrit ensuite la réalisation du pipeline, l'évaluation du prototype, les contrôles de sécurité et les limites à lever avant une mise en production.",
        "The report first analyzes the institutional, regulatory and operational context, then presents the functional and technical architecture of the solution. It highlights how the implementation choices answer data sovereignty, auditability, document security and customer experience requirements.": "The report first analyzes the institutional, regulatory and operational context, then presents the requirements, data flow and processing architecture. It then describes the implementation, prototype evaluation, security controls and remaining limits before production use.",
        "BICEC VeriPass est conçu comme un pipeline d'acquisition client et de gestion de preuves numériques. La solution se compose d'une interface client PWA, d'un back office de validation, d'une API FastAPI, de workers Celery/Redis pour les traitements asynchrones, d'une base PostgreSQL et d'un stockage documentaire local.": "BICEC VeriPass est conçu comme un pipeline d'onboarding KYC et de gestion de preuves numériques. La solution se compose d'une interface client PWA, d'un back office de validation, d'une API FastAPI, de workers Celery/Redis pour les traitements asynchrones, d'une base PostgreSQL et d'un stockage documentaire local.",
        "Le principe directeur est celui de l'automatisation contrôlée : les moteurs OCR et biométriques préparent le dossier, calculent des scores et signalent les anomalies, mais l'agent habilité conserve la responsabilité de la validation finale. VeriPass agit donc comme un outil d'aide à l'instruction KYC, sans intégrer directement les bases transactionnelles de production.": "Les moteurs OCR et biométriques préparent le dossier, calculent des scores et signalent les anomalies, mais l'agent habilité conserve la responsabilité de la validation finale. VeriPass agit donc comme un outil d'aide à l'instruction KYC, sans intégrer directement les bases transactionnelles de production.",
        "Ce premier chapitre a établi la problématique réglementaire, commerciale, technique et économique de VeriPass. La solution doit réduire la friction de l'ouverture de compte tout en respectant les exigences de la zone CEMAC. L'architecture proposée répond à cette contrainte par un pipeline numérique structuré, auditable et compatible avec une validation humaine. Le chapitre suivant formalise le besoin, les données d'identité, les exigences fonctionnelles et non fonctionnelles, puis le cadrage méthodologique du projet.": "Ce premier chapitre a établi la problématique réglementaire, commerciale, technique et économique de VeriPass. La solution doit réduire les reprises de saisie, les relances et les délais d'instruction, tout en respectant les exigences de la zone CEMAC. L'architecture proposée répond à cette contrainte par un pipeline numérique structuré, auditable et compatible avec une validation humaine. Le chapitre suivant formalise le besoin, les données d'identité, les exigences fonctionnelles et non fonctionnelles, puis le cadrage méthodologique du projet.",
        "Le reverse proxy reverse joue le rôle de point d'entrée TLS et distribue les requêtes vers la PWA mobile, le back office et l'API. Cette séparation permet de protéger les interfaces internes, de centraliser certains en têtes de sécurité et d'isoler les temps longs liés à la capture documentaire.": "Le proxy inverse Nginx joue le rôle de point d'entrée TLS et distribue les requêtes vers la PWA mobile, le back office et l'API. Cette séparation protège les interfaces internes, centralise certains en têtes de sécurité et isole les temps longs liés à la capture documentaire.",
        "La trajectoire fonctionnelle peut être résumée par les états suivants : DRAFT, SUBMITTED, OCR_PROCESSING, BIOMETRIC_PROCESSING, READY_FOR_REVIEW, puis APPROVED ou REJECTED. Dans le code actuel, certains libellés opérationnels diffèrent, notamment PENDING_AGENT_REVIEW, PENDING_INFO ou FRAUD_SUSPECT. Le principe reste identique : séparer la préparation du dossier, les traitements automatiques et la décision humaine.": "La trajectoire réellement utilisée par le prototype distingue la préparation, la revue et la décision. Le dossier démarre en DRAFT, passe en PENDING_AGENT_REVIEW après soumission, peut revenir en PENDING_INFO lorsqu'un complément est demandé, puis se termine par APPROVED, REJECTED ou FRAUD_SUSPECT selon la décision back office. Cette machine d'états évite de confondre un dossier encore éditable, un dossier en attente de revue et un dossier déjà décidé.",
        "Ce deuxième chapitre a traduit le cadre d'onboarding en exigences logicielles formelles et a présenté notre conduite de projet agile et hybride basée sur le Cycle en V et Scrum. Nous avons étayé l'état de l'art scientifique de l'OCR et de la biométrie faciale et formalisé les exigences non fonctionnelles de performance et de calibrage sous contraintes matérielles réelles. Enfin, nous avons dressé le bilan de nos acquis d'apprentissage. Le chapitre suivant détaille l'architecture technique globale, les conteneurs Docker et la modélisation de la base de données relationnelle répondant à ces exigences.": "",
    }
    for old, new in replacements.items():
        if new:
            replace_paragraph(doc, old, new)
        else:
            for p in list(doc.paragraphs):
                if p.text.strip() == old:
                    remove_block(p)

    # Remove duplicate synthetic tables and table-heavy blocks that weakened the thesis voice.
    remove_tables_by_header(
        doc,
        (
            "Partie |",
            "Elément | Information",
            "Exigence KYC |",
            "Limite observée |",
            "Composant de coût |",
            "Acteur |",
            "Rôle |",
            "Famille de données |",
            "Domaine |",
            "Domaine d'exigence |",
            "Phase |",
            "Phase de stage |",
            "Objectif |",
        ),
    )
    remove_para_with_following_source(
        doc,
        (
            "Tableau 5. Impacts économiques attendus",
            "Tableau 6. Acteurs du système",
            "Tableau 7. Synthèse des exigences",
            "Tableau 9. Planning",
            "Tableau 12. Bilan",
        ),
    )
    for p in list(doc.paragraphs):
        source = p.text.strip()
        if source in {
            "Source : synthèse de l'auteur, d'après le rapport de cadrage financier BICEC VeriPass.",
            "Source : synthèse de l'auteur, d'après les rôles applicatifs et la structure RBAC du code backend.",
            "Source : synthèse de l'auteur, d'après les spécifications d'ingénierie et le plan de tests VeriPass.",
            "Source : synthèse de l'auteur, d'après le calendrier de stage et la roadmap agile VeriPass.",
        }:
            remove_block(p)
    for p in list(doc.paragraphs):
        if p.text.strip() in {"2.11 Synthèse du chapitre"}:
            remove_block(p)
        if p.text.strip() in {"Tableau 1. Structure du mémoire", "Tableau 1. Structure prévue du mémoire"}:
            remove_block(p)

    for table in doc.tables:
        head = table_header(table)
        if head.startswith("Limite constatée |"):
            for row in table.rows:
                if row.cells and row.cells[0].text.strip() == "Friction physique":
                    row.cells[1].text = "Délais de constitution variables, dépendants des déplacements, de la complétude du dossier et des relances"
                    row.cells[2].text = "Parcours mobile visant à réduire les reprises et les relances, à mesurer en phase pilote"

    # Update preliminary lists after removing redundant tables.
    list_replacements = {
        "Tableau 5. Limites du processus manuel et réponse attendue du pipeline numérique": "Tableau 5. Architecture logique et responsabilités techniques",
        "Tableau 6. Acteurs du système et responsabilités attendues": "Tableau 6. Gestion des échecs OCR et biométriques",
        "Tableau 7. Données KYC manipulées et critères de qualité": "",
        "Tableau 8. Exigences, contraintes, risques et réponses méthodologiques": "",
        "Tableau 9. Planning de conduite, risques et modes de collaboration": "",
        "Tableau 10. Architecture logique et responsabilités techniques": "",
        "Tableau 11. Gestion des échecs OCR et biométriques": "",
        "Tableau 12. Bilan synthétique des objectifs et des réalisations": "",
        "Figure 1. Chaîne fonctionnelle du pipeline KYC VeriPass": "Figure 1. Chaîne fonctionnelle du pipeline KYC VeriPass\nFigure 2. Architecture logique simplifiée de BICEC VeriPass\nFigure 3. Pipeline de données et états principaux du dossier KYC\nFigure 4. MLD simplifié des entités KYC principales",
        "Tableau 10. Architecture logique et responsabilités techniques": "Tableau 5. Architecture logique et responsabilités techniques",
        "Tableau 11. Gestion des échecs OCR et biométriques": "Tableau 6. Gestion des échecs OCR et biométriques",
    }
    for p in list(doc.paragraphs):
        t = p.text.strip()
        if t in list_replacements:
            if list_replacements[t]:
                set_text(p, list_replacements[t])
            else:
                remove_block(p)

    # Strengthen chapter 2 with prose rather than new tables.
    target = "La qualité des données doit être vérifiée le plus tôt possible. Côté client, les consignes de capture et les contrôles visuels réduisent les images floues, mal cadrées ou partiellement lisibles. Côté serveur, le calcul d'une empreinte SHA-256 lors de l'upload permet de contrôler l'intégrité des pièces et de relier chaque fichier à son dossier."
    for p in doc.paragraphs:
        if p.text.strip() == target:
            cur = p
            cur = insert_after(cur, "Le lignage des données est un point structurant du besoin. Une information saisie ou extraite doit pouvoir être suivie depuis la capture mobile jusqu'à la décision back office. Le dossier relie donc la session client, le document stocké, le champ OCR extrait, la correction éventuelle, le score biométrique, la décision et le journal d'audit. Cette chaîne limite les pertes de contexte et rend les contrôles ultérieurs plus simples à reconstruire.", "Memoire Body")
            insert_after(cur, "La gouvernance des données doit rester proportionnée à la sensibilité du KYC. Les pièces d'identité, les images de visage et les journaux de décision ne sont pas de simples fichiers techniques. Ils portent une finalité réglementaire, doivent être accessibles seulement aux rôles habilités, et doivent rester conservés avec une trace suffisante pour justifier les décisions prises.", "Memoire Body")
            break

    # Replace state of the art with verified, calibrated claims.
    old_131 = "Le développement d'un pipeline KYC impose de s'appuyer sur la littérature relative à l'extraction documentaire. Les CNI photographiées depuis un téléphone présentent des difficultés connues : usure du support, reflets, compression, cadrage imparfait, fonds complexes et variation des polices. Les travaux de Nguyen, Nguyen et Nguyen (2020) ainsi que ceux de Raj, Sreenivas et Jawahar (2020) montrent que l'extraction documentaire gagne en robustesse lorsque le pipeline combine localisation des zones utiles et reconnaissance textuelle. Dans VeriPass, cette logique justifie l'usage d'un moteur OCR rapide pour les cas simples, avec un traitement asynchrone plus poussé lorsque la capture est difficile ou incomplète."
    new_131 = "Le développement d'un pipeline KYC impose de s'appuyer sur la littérature relative à l'extraction documentaire. Les CNI photographiées depuis un téléphone présentent des difficultés connues : usure du support, reflets, compression, cadrage imparfait, fonds complexes et variation des polices. Les travaux sur la reconnaissance des documents d'identité rappellent que ce problème dépasse l'OCR classique, car le système doit extraire des champs tout en contribuant à la vérification d'identité et à la prévention de la fraude (Bulatov et al., 2022)."
    replace_paragraph(doc, old_131, new_131)
    for p in doc.paragraphs:
        if p.text.strip() == new_131:
            cur = p
            cur = insert_after(cur, "Dans ce contexte, l'approche retenue par VeriPass reste volontairement hybride. PaddleOCR est utilisé comme première passe locale pour détecter et reconnaître le texte, dans l'esprit des systèmes OCR légers conçus pour équilibrer vitesse et précision (Du et al., 2020). Lorsque la qualité de capture ou la structure du document rend cette première passe insuffisante, le pipeline prévoit un traitement différé et une revue humaine, plutôt qu'une acceptation silencieuse d'un champ incertain.", "Memoire Body")
            cur = insert_after(cur, "Les recherches récentes sur l'extraction d'informations dans les documents d'identité montrent aussi la rareté des jeux de données réellement exploitables, pour des raisons de confidentialité et de sécurité (Carta et al., 2024). Cette contrainte justifie, dans le cadre du mémoire, de distinguer ce qui est implémenté, ce qui est mesuré sur les preuves disponibles, et ce qui devra être calibré lors d'un pilote plus large.", "Memoire Body")
            break
    old_132 = "Pour la vérification d'identité faciale, la reconnaissance du visage ne suffit pas. Les attaques par présentation, notamment photo, vidéo ou masque, obligent à combiner comparaison faciale et détection de vivacité. Yu et al. (2020) soulignent l'intérêt de coupler le face matching avec des tests de liveness. Dans VeriPass, cette orientation se traduit par un challenge actif guidé par l'interface, puis par une comparaison entre le selfie et la photo de la CNI. Les résultats doivent rester exprimés en scores, seuils et motifs afin de permettre une revue humaine en cas d'incertitude."
    new_132 = "Pour la vérification d'identité faciale, la reconnaissance du visage ne suffit pas. Les attaques par présentation, notamment photo, vidéo ou masque, obligent à combiner comparaison faciale et détection de vivacité. Les revues récentes sur le face anti spoofing soulignent que les méthodes doivent être évaluées face à des conditions de capture et à des attaques variées (Yu et al., 2023). Dans VeriPass, cette orientation se traduit par un challenge actif guidé par l'interface, puis par une comparaison entre le selfie et la photo de la CNI. Les résultats restent exprimés en scores, seuils et motifs afin de permettre une revue humaine en cas d'incertitude."
    replace_paragraph(doc, old_132, new_132)

    for p in doc.paragraphs:
        if p.text.strip() == "Les exigences fonctionnelles cadrent les flux applicatifs du système : authentification, capture guidée, preuve de vie, confirmation des champs OCR, collecte de l'adresse, gestion du NIU, consentement, signature, soumission du dossier et revue back office. Elles prévoient aussi les cas dégradés : échec de capture, document illisible, OCR incertain, liveness non concluant, demande de complément et décision motivée.":
            insert_after(p, "Pour conserver une traçabilité d'ingénierie, chaque exigence doit pouvoir être reliée à un composant et à une preuve. La capture CNI renvoie aux endpoints de capture et au stockage documentaire; la revue OCR renvoie aux champs extraits et corrigés; la preuve de vie renvoie au challenge mobile et au résultat serveur; la décision renvoie au back office et au journal d'audit. Cette lecture évite de traiter les exigences comme une liste abstraite.", "Memoire Body")
            break

    for p in doc.paragraphs:
        if p.text.strip() == "Les exigences non fonctionnelles portent sur la sécurité, la performance, l'auditabilité, la maintenabilité et l'expérience utilisateur. Les traitements OCR et biométriques doivent rester compatibles avec un parcours fluide. Les pièces et métadonnées doivent être protégées, les accès limités par rôle, les décisions tracées et les erreurs rendues compréhensibles pour le client comme pour les agents.":
            insert_after(p, "La performance est donc traitée comme une contrainte de parcours, non comme une promesse isolée. Les traitements longs doivent être asynchrones, les statuts doivent rester lisibles pour l'utilisateur, et les erreurs doivent ouvrir une reprise contrôlée. La sécurité, de son côté, repose sur la séparation des rôles, l'authentification, la journalisation, l'intégrité documentaire et la limitation des accès aux données sensibles.", "Memoire Body")
            break

    # Chapter 4 and chapter 5 full prose draft.
    ch4_sections = [
        (
            "4.1 Environnement de réalisation et organisation du dépôt",
            [
                "La réalisation de VeriPass s'appuie sur un environnement reproductible afin de limiter les écarts entre conception, développement et démonstration. Le dépôt est organisé autour de trois surfaces principales : le backend FastAPI, l'application mobile PWA et le back office. À ces trois surfaces s'ajoutent les fichiers d'infrastructure, la documentation d'architecture, les contrats API, les modèles de données et les preuves de tests.",
                "L'environnement Docker Compose regroupe les services nécessaires au prototype : Nginx comme point d'entrée TLS, la PWA mobile, le back office, l'API, PostgreSQL, Redis, les workers Celery, les tâches planifiées et Flower pour la supervision des files. Cette organisation ne constitue pas encore une architecture de production bancaire, mais elle rend le prototype exécutable, observable et démontrable sur une machine de développement.",
                "Le choix d'un environnement conteneurisé répond aussi à une exigence pédagogique et industrielle. Il permet de montrer l'ensemble de la chaîne sans dépendre d'un service externe pour les données sensibles. Les volumes documentaires, les volumes de base et les volumes de sauvegarde matérialisent la séparation entre code applicatif, métadonnées relationnelles et fichiers KYC.",
            ],
        ),
        (
            "4.2 Réalisation du backend FastAPI",
            [
                "Le backend porte les règles de domaine et les contrats d'échange. Il ne se limite pas à recevoir des formulaires : il crée ou reprend une session KYC, accepte les pièces, calcule les métadonnées documentaires, déclenche l'OCR, reçoit les corrections, enregistre le consentement, vérifie la capacité de soumission et transmet le dossier au back office.",
                "Le module KYC expose les routes principales du parcours : démarrage de session, capture CNI, lecture des champs OCR, revue ou confirmation des champs, soumission de l'adresse, soumission du NIU, consentement, signature, liveness et soumission finale. Cette granularité donne une bonne maîtrise du parcours, car chaque étape peut échouer, être reprise ou être auditée séparément.",
                "Les contrats d'API jouent un rôle méthodologique important. Ils forcent à préciser ce qu'un écran peut envoyer, ce que le backend accepte, et ce que le client peut attendre en retour. Dans un projet KYC, cette précision évite les ambiguïtés entre document capturé, document exploitable, champ extrait, champ confirmé et dossier soumis.",
                "La logique de readiness illustre cette discipline. Avant la soumission, l'API vérifie que les éléments requis sont présents. Le but n'est pas seulement de bloquer un formulaire incomplet; il s'agit d'empêcher l'entrée dans la file back office d'un dossier qui ne peut pas être instruit correctement. Cette règle économise du temps agent et améliore la qualité du flux.",
                "Le backend applique aussi les transitions d'état. Un dossier en préparation reste éditable, tandis qu'un dossier soumis passe en attente de revue. Une demande de complément ramène le client vers une reprise limitée, sans effacer l'historique. Une décision approuvée, rejetée ou marquée comme suspecte clôt le cycle métier et modifie le niveau d'accès associé.",
                "La sécurité applicative repose sur l'authentification, la séparation client et agent, les rôles back office et la journalisation des actions sensibles. Les rôles applicatifs ne sont pas des personnages fictifs : ils correspondent aux responsabilités fonctionnelles du processus, notamment agent KYC ou chargé de clientèle, analyste AML/CFT, responsable conformité et administrateur IT.",
            ],
        ),
        (
            "4.3 Réalisation du parcours mobile",
            [
                "La PWA mobile porte le parcours client. Elle guide l'utilisateur depuis l'authentification jusqu'à la soumission du dossier. Les étapes principales sont l'authentification par OTP, la capture de la CNI recto verso, la vérification des champs extraits, la preuve de vie, la collecte de l'adresse, le consentement, la signature et l'envoi du dossier au back office.",
                "Le parcours a été conçu pour rester séquentiel. Cette décision évite de demander au client de comprendre la logique interne du KYC. Elle permet aussi de contrôler la complétude avant soumission : un dossier ne doit pas être transmis si les pièces attendues, le consentement ou les étapes de vérification essentielles sont absents.",
                "La reprise de parcours constitue un point important de l'expérience. Dans un environnement mobile, l'utilisateur peut interrompre la procédure, perdre la connexion ou devoir reprendre une capture. Les états de session et le champ de dernière étape permettent de ramener l'utilisateur au bon endroit, sans recommencer inutilement tout le dossier.",
                "Le parcours mobile tient aussi compte de la contrainte réseau. Une expérience KYC ne peut pas supposer une connexion parfaite, surtout lorsque la capture de pièces et de selfies produit des fichiers plus lourds qu'un formulaire classique. La logique de reprise et de synchronisation permet de réduire les abandons liés aux interruptions, tout en évitant que des étapes dépendantes soient envoyées dans le mauvais ordre.",
                "La gestion de l'appareil participe à la sécurité du parcours. Lorsqu'un appareil est enregistré, certaines actions sensibles peuvent exiger un identifiant cohérent. Ce mécanisme ne remplace pas l'authentification, mais il ajoute un signal de contrôle utile pour limiter les usages anormaux d'un compte client.",
                "Le mobile ne prend pas la décision KYC. Il collecte les preuves et rend les corrections possibles. Cette limite est importante : le client peut confirmer ou corriger des champs OCR, mais la décision bancaire reste dans le back office, avec un agent habilité et une trace exploitable.",
            ],
        ),
        (
            "4.4 Intégration du pipeline OCR",
            [
                "L'intégration OCR commence au moment de l'upload documentaire. Chaque pièce est enregistrée avec son type, son chemin relatif, sa taille, son hash et son statut de traitement. Ce choix permet de relier un résultat OCR à un fichier précis, au lieu de manipuler des textes extraits sans preuve documentaire.",
                "PaddleOCR sert de première passe locale. Son intérêt est de fournir un résultat rapide sans envoyer les pièces d'identité vers un service externe. Le pipeline applique ensuite des règles de post traitement adaptées aux champs attendus d'une CNI : noms, prénoms, dates, numéro de document, poste d'identification, profession ou adresse selon le côté de la pièce et la lisibilité de la capture.",
                "Les résultats OCR ne sont pas considérés comme une vérité définitive. Les champs sont conservés avec un score de confiance et un statut. Lorsque l'extraction est incomplète ou incertaine, le système peut marquer le document comme partiel, demander une correction ou orienter vers un traitement différé. Cette gestion explicite de l'incertitude est plus sûre qu'une extraction affichée comme certaine.",
                "La correction humaine est intégrée au pipeline. Le client peut confirmer ou corriger les champs dans le parcours, et le back office peut intervenir sur les champs lors de la revue. La valeur initiale, la valeur corrigée, l'auteur et l'horodatage doivent rester disponibles pour comprendre l'écart entre l'extraction automatique et la décision finale.",
                "Le traitement GLM-OCR est prévu comme mécanisme de renfort pour les cas plus complexes. Dans le mémoire, il doit être présenté avec prudence : il complète la stratégie d'extraction lorsque la première passe OCR ne suffit pas, mais il ne remplace ni la validation humaine ni la nécessité de mesurer les résultats sur un jeu représentatif.",
            ],
        ),
        (
            "4.5 Intégration de la vérification biométrique",
            [
                "La vérification biométrique a été intégrée comme une chaîne en deux temps. Le premier temps concerne la preuve de vie. Le client suit un challenge actif, par exemple sourire, cligner des yeux ou tourner la tête. Le second temps concerne la comparaison faciale entre le selfie et la photo portée par la pièce d'identité.",
                "Côté mobile, MediaPipe permet d'exploiter les repères du visage pendant le challenge. Le but n'est pas de transformer le téléphone en moteur de décision, mais de guider l'utilisateur et de produire une preuve technique exploitable par le backend. Cette séparation limite la dépendance à l'appareil du client et garde la décision centralisée.",
                "Côté backend, le résultat de liveness est enregistré avec un statut, un score et un motif. Lorsque les dépendances de comparaison faciale sont disponibles, DeepFace permet de produire un statut de face matching, une distance et un seuil. Ces informations sont utiles à la revue, mais elles ne doivent pas être interprétées comme une décision bancaire automatique.",
                "Les échecs biométriques sont traités comme des cas métier. Une mauvaise lumière, une caméra faible, un visage mal cadré ou une consigne mal comprise peuvent produire un résultat négatif sans fraude. Le système doit donc prévoir des reprises limitées et une escalade humaine lorsque le score ne permet pas de conclure proprement.",
            ],
        ),
        (
            "4.6 Réalisation du back office et des décisions",
            [
                "Le back office transforme les données collectées en dossier instruisable. Il présente les pièces, les champs extraits, les corrections, les scores, les alertes et l'historique. Cette interface répond à une logique différente du mobile : elle ne cherche pas à collecter, elle cherche à décider avec suffisamment de preuves.",
                "La file de dossiers organise la revue des demandes soumises. L'agent KYC ou chargé de clientèle peut ouvrir un dossier, consulter les documents, demander un complément, approuver ou rejeter. L'analyste AML/CFT intervient sur les alertes et les suspicions. La responsable conformité supervise les décisions sensibles et les indicateurs de contrôle.",
                "Les demandes de complément constituent un cas important du parcours. Elles évitent de rejeter un client uniquement parce qu'une pièce est incomplète ou difficile à lire. Le back office peut demander une information complémentaire, et le client peut la fournir sans redémarrer tout le processus. Cette boucle améliore l'expérience client et maintient la qualité du dossier.",
                "L'intégration AML/CFT donne au back office une profondeur supplémentaire. Les alertes de sanctions, les conflits de NIU, les documents expirants ou les suspicions de doublon ne relèvent pas du même traitement qu'une simple erreur de saisie. VeriPass prépare donc des signaux distincts afin que chaque rôle puisse traiter le dossier selon sa responsabilité.",
                "Chaque décision doit être motivée. Cette exigence est centrale pour la conformité. Un rejet, une demande de complément ou un signalement de fraude ne doit pas être une action muette. Le système doit conserver l'acteur, le motif, la date, l'adresse IP lorsque disponible et les données ayant soutenu la décision.",
                "L'audit log complète cette logique. Il donne une trace transversale des actions sensibles : revue, affectation, classification, correction, décision, sauvegarde ou export. Cette trace ne remplace pas la conformité métier, mais elle fournit un support technique à la reconstitution des dossiers.",
            ],
        ),
        (
            "4.7 Difficultés rencontrées et arbitrages",
            [
                "La première difficulté concerne la qualité des captures. Les CNI photographiées par téléphone peuvent être floues, mal cadrées, compressées ou partiellement masquées par des reflets. Cette contrainte a conduit à traiter l'OCR comme un pipeline incertain, avec scores, statuts, reprises et corrections, plutôt que comme une extraction définitive.",
                "La deuxième difficulté concerne les traitements longs. L'OCR et le liveness peuvent dépasser le temps acceptable d'un appel interactif. L'usage de Redis et Celery répond à cette contrainte en séparant la requête client du traitement différé. Le client ne doit pas attendre sans explication; le dossier doit porter un statut lisible.",
                "Le choix asynchrone a un coût de conception. Il oblige à penser les statuts, les reprises et les messages d'attente avec rigueur. Un traitement en arrière plan mal expliqué crée de l'incertitude pour le client et pour l'agent. Le prototype traite donc le statut du dossier comme une information métier, et non comme un détail technique.",
                "La troisième difficulté concerne le niveau de maturité biométrique. Un score de comparaison faciale dépend du modèle, de la qualité de la photo CNI, de la lumière et des seuils retenus. Le projet ne prétend donc pas que la biométrie suffit à décider. Il l'utilise comme signal de revue, avec escalade humaine lorsque le résultat est incertain.",
                "La quatrième difficulté concerne la frontière entre prototype et système bancaire. Plusieurs composants démontrent le flux, mais certains choix doivent encore être validés par les équipes sécurité, conformité et exploitation. Cette distinction est assumée dans le mémoire afin de ne pas présenter une démonstration de fin d'études comme une solution déjà certifiée.",
                "Le dernier arbitrage concerne l'intégration bancaire. Le prototype n'intègre pas directement un core banking de production. Ce choix évite de confondre preuve de concept KYC et activation commerciale réelle. Il permet de concentrer le mémoire sur la chaîne d'acquisition, de vérification, de traçabilité et de revue.",
            ],
        ),
        (
            "4.8 Synthèse du chapitre",
            [
                "Ce chapitre a présenté la réalisation concrète de VeriPass. Le système associe une PWA de collecte, une API FastAPI, des traitements asynchrones, une base relationnelle, un stockage documentaire, un pipeline OCR, une vérification biométrique et un back office de décision. L'ensemble forme un prototype cohérent d'onboarding KYC, sans prétendre remplacer les validations réglementaires et opérationnelles nécessaires à une mise en production.",
                "L'apport d'ingénierie du chapitre réside surtout dans l'assemblage contrôlé des composants. Une application mobile seule n'aurait pas suffi; un moteur OCR isolé n'aurait pas répondu à la conformité; une base de données sans journalisation n'aurait pas permis l'audit. La valeur de VeriPass vient de la liaison entre ces briques : chaque capture produit une preuve, chaque preuve alimente un traitement, chaque traitement produit un statut, et chaque statut prépare une décision traçable.",
                "La réalisation confirme le principe directeur du mémoire : automatiser les tâches répétitives et structurables, tout en conservant la responsabilité humaine sur les décisions sensibles. Le chapitre suivant évalue ce prototype, précise les preuves disponibles, puis discute la sécurité, le déploiement, l'impact et les limites.",
            ],
        ),
    ]
    replace_block_between(
        doc,
        "Chapitre 4 : Réalisation et intégration des mécanismes intelligents de vérification",
        "Chapitre 5 : Évaluation, sécurité, déploiement, impact et limites",
        ch4_sections,
    )

    ch5_sections = [
        (
            "5.1 Périmètre et protocole d'évaluation",
            [
                "L'évaluation porte sur un prototype avancé, exécuté en environnement local conteneurisé. Elle ne doit pas être lue comme une certification de production bancaire. Son objectif est de vérifier que la chaîne fonctionnelle peut être exécutée de bout en bout, que les principaux statuts sont cohérents, que les traitements documentaires produisent des résultats exploitables et que les limites restantes sont identifiées.",
                "Le protocole combine trois familles de preuves. La première concerne le parcours API complet, depuis l'authentification jusqu'à l'approbation back office. La deuxième concerne la boucle OCR sur des images de CNI utilisées pour valider le comportement réel du pipeline. La troisième concerne les contrôles de sécurité, de rôle, d'audit et de déploiement observables dans le dépôt.",
                "L'évaluation suit une logique de traçabilité. Chaque preuve doit être reliée à une exigence : la soumission vérifie la complétude du parcours, l'OCR vérifie l'extraction documentaire, le liveness vérifie l'intégration biométrique, le back office vérifie la décision humaine, et l'audit log vérifie la reconstitution des actions. Cette lecture évite de présenter des tests isolés sans rapport avec la problématique.",
                "Ce choix d'évaluation est volontairement prudent. Les résultats disponibles démontrent la faisabilité technique du prototype, mais ils ne suffisent pas à conclure sur des taux de précision réglementaires ou biométriques en production. Une phase pilote avec données réelles contrôlées restera nécessaire.",
            ],
        ),
        (
            "5.2 Parcours fonctionnel de bout en bout",
            [
                "Le scénario de bout en bout documenté dans les preuves de test exécute les étapes principales du parcours : envoi et vérification OTP, enregistrement de l'appareil, démarrage de session KYC, upload des pièces CNI, upload d'un justificatif, upload du selfie, revue OCR, liveness, adresse, consentement, NIU, signature, vérification de readiness, soumission, traitement back office, demande d'information, ajout d'un complément et approbation.",
                "Les appels API du scénario retournent des statuts HTTP cohérents. Les opérations de capture CNI sont les plus longues, avec des durées de l'ordre de plusieurs secondes en raison du traitement OCR. Les actions purement transactionnelles, comme l'adresse, le consentement, le NIU, la signature ou la soumission, restent beaucoup plus courtes. Cette différence confirme l'intérêt de séparer les traitements lourds et les transitions métier.",
                "Le résultat fonctionnel est significatif : le dossier peut être soumis, passer en PENDING_AGENT_REVIEW, revenir en PENDING_INFO lorsqu'un complément est demandé, puis être approuvé. Le client voit l'évolution de son statut et le back office dispose d'un dossier exploitable. Ce scénario donne une preuve concrète de l'enchaînement entre mobile, API, base, support et décision.",
                "Ce scénario est également utile pour vérifier la cohérence des rôles. Le client ne peut pas décider sur son propre dossier, l'agent ne modifie pas les étapes de capture à la place du client, et les décisions back office restent séparées des scores automatiques. Cette séparation donne une lecture claire des responsabilités.",
            ],
        ),
        (
            "5.3 Évaluation du traitement OCR",
            [
                "La boucle OCR documentée dans le dépôt a porté sur un lot de 60 identifiants CNI. Le rapport indique qu'au terme des passes initiales et des reprises ciblées, chaque identifiant CNI a obtenu au moins une passe OCR réelle réussie dans le parcours applicatif. Ce résultat est utile, car il montre que le pipeline fonctionne sur un ensemble varié de captures, avec interaction réelle avec les endpoints et les données de revue.",
                "Ce résultat ne doit toutefois pas être surinterprété. Il ne constitue pas encore une mesure de précision par champ, ni un taux d'erreur statistiquement représentatif de toutes les CNI camerounaises rencontrées sur le terrain. Il valide plutôt la capacité du pipeline à produire une revue OCR exploitable, à gérer des reprises et à transformer des observations manuelles en règles de post traitement.",
                "Les corrections apportées dans la boucle OCR sont instructives. Certaines concernaient l'adresse, la profession, les dates, le numéro CNI ou le poste d'identification. Elles montrent que le travail d'ingénierie ne se limite pas à appeler un moteur OCR : il faut normaliser les formats, protéger les valeurs légitimes, tenir compte des zones du document et conserver l'incertitude lorsque la capture reste ambiguë.",
            ],
        ),
        (
            "5.4 Évaluation biométrique et limites des scores",
            [
                "Le scénario de bout en bout montre que la soumission liveness peut être exécutée et que le système renvoie un résultat de vivacité. Cette preuve valide l'intégration fonctionnelle du challenge, du selfie et du backend. Elle ne suffit pas encore à établir un taux de réussite biométrique ou un taux de faux rejet, car ces métriques exigent un jeu de tests représentatif et contrôlé.",
                "La biométrie doit donc rester cadrée comme un signal d'aide à la revue. Le score de liveness et le statut de face matching donnent au back office une information supplémentaire, mais ne remplacent pas la décision humaine. Cette position réduit le risque de rejet injustifié et maintient la responsabilité métier au niveau des acteurs habilités.",
                "Les limites connues concernent la lumière, la qualité de la caméra, la différence entre la photo CNI et le visage actuel, les attaques par présentation et les biais possibles des modèles. Avant une mise en production, les seuils devront être calibrés avec des données conformes aux règles de protection des données personnelles et validés avec les équipes conformité.",
            ],
        ),
        (
            "5.5 Sécurité, conformité et auditabilité",
            [
                "La sécurité du prototype repose d'abord sur la séparation des rôles. Le client, l'agent KYC, l'analyste AML/CFT, la responsable conformité et l'administrateur IT ne disposent pas des mêmes droits. Cette séparation réduit le risque qu'un acteur manipule un dossier en dehors de son périmètre fonctionnel.",
                "L'authentification, les tokens, les rôles back office et le contrôle d'appareil structurent l'accès aux parcours sensibles. Les documents sont conservés avec des métadonnées et des empreintes, ce qui permet de relier un fichier à un dossier et de détecter certaines incohérences. Les décisions importantes doivent être journalisées avec l'acteur, l'action, le moment et le motif.",
                "L'auditabilité répond directement aux exigences KYC. Un contrôle interne ou externe doit pouvoir comprendre comment un dossier a été constitué, quelles pièces ont été utilisées, quels champs ont été extraits ou corrigés, quel score a été obtenu et qui a pris la décision. La base relationnelle et l'audit log fournissent le support technique de cette reconstitution.",
                "La protection des données biométriques demande une attention particulière. Le selfie et les résultats de comparaison faciale ne sont pas des données ordinaires. Leur usage doit être justifié par la finalité KYC, limité aux acteurs habilités, conservé selon une politique documentée et supprimé ou archivé selon les règles applicables.",
                "La conformité ne se démontre donc pas par une seule fonctionnalité. Elle résulte d'un ensemble : consentement, minimisation, contrôle d'accès, conservation des preuves, journalisation, revue humaine, capacité d'export et procédure de gestion des incidents. VeriPass prépare ces mécanismes, mais leur validation finale relève d'une gouvernance bancaire.",
                "Ces contrôles ne suffisent pas encore à déclarer une conformité de production. Une industrialisation demanderait une revue complète des secrets, une politique de chiffrement au repos adaptée à l'hébergement, une procédure de sauvegarde et restauration testée, une analyse de risques, des tests d'intrusion et une validation juridique du traitement des données biométriques.",
            ],
        ),
        (
            "5.6 Déploiement, exploitation et observabilité",
            [
                "Le déploiement local repose sur Docker Compose. Cette approche permet de démarrer l'API, les interfaces, PostgreSQL, Redis, les workers et les services de supervision dans un environnement cohérent. Elle facilite la démonstration et la reprise par un autre développeur, car les dépendances principales sont décrites dans un même ensemble.",
                "L'exploitation d'un tel système exige toutefois plus qu'un démarrage de conteneurs. Il faut surveiller la disponibilité de l'API, la santé de la base, la saturation des files Celery, la taille des volumes documentaires, les erreurs OCR, les temps de traitement et l'état des sauvegardes. Le prototype prépare cette logique par ses journaux, ses healthchecks et ses documents d'architecture.",
                "Les workers constituent un point sensible de l'exploitation. Si la file OCR ralentit, l'application peut rester accessible tout en bloquant la progression des dossiers. La supervision doit donc distinguer disponibilité web, disponibilité de la base, capacité des files et état des traitements différés. Cette distinction est indispensable pour diagnostiquer rapidement une panne partielle.",
                "En production, la supervision devrait être complétée par des alertes, des tableaux de bord d'exploitation, une politique de rotation des journaux et une stratégie de restauration. Les traitements OCR et biométriques doivent aussi être surveillés comme des composants métier, car leur dégradation peut bloquer l'entrée en relation sans rendre l'application totalement indisponible.",
            ],
        ),
        (
            "5.7 Impact attendu pour la BICEC",
            [
                "L'impact principal de VeriPass est opérationnel. Le système vise à réduire les reprises de saisie, à rendre les pièces plus faciles à contrôler, à centraliser les preuves et à donner au back office un dossier structuré. Ces gains ne doivent pas être exprimés comme des économies définitives sans pilote, mais comme des leviers mesurables lors d'une expérimentation terrain.",
                "Pour le client, l'impact attendu concerne la simplicité du parcours et la possibilité de commencer la constitution du dossier à distance. Pour les agents, l'intérêt porte sur la réduction des dossiers incomplets, la disponibilité des champs extraits, la meilleure lisibilité des anomalies et la conservation de l'historique. Pour la conformité, l'apport principal réside dans la traçabilité et la capacité de reconstituer la décision.",
                "Pour la direction, l'intérêt porte sur la mesure. Un processus papier rend plus difficile le suivi des reprises, des temps d'attente, des motifs de rejet et de la charge par rôle. Un parcours numérique permet de produire des indicateurs plus fiables, à condition que les événements soient bien journalisés et interprétés avec prudence.",
                "L'impact humain est également important. Le projet montre que l'automatisation bancaire ne consiste pas à remplacer les acteurs, mais à déplacer leur effort vers les tâches à plus forte valeur : arbitrer, contrôler, expliquer, traiter les exceptions et améliorer le dispositif.",
            ],
        ),
        (
            "5.8 Limites et perspectives d'industrialisation",
            [
                "La première limite est l'absence de pilote de production sur un échantillon client réel et représentatif. Les preuves disponibles valident le prototype et certains scénarios, mais elles ne remplacent pas des mesures terrain de temps de parcours, de taux d'abandon, de précision OCR, de reprise documentaire et de satisfaction des agents.",
                "La deuxième limite concerne la calibration des seuils. Les seuils OCR, liveness et face matching doivent être ajustés sur des données collectées légalement, annotées correctement et représentatives des conditions de capture locales. Un seuil trop strict peut rejeter des clients légitimes; un seuil trop permissif peut laisser passer des erreurs ou des risques de fraude.",
                "La troisième limite concerne l'intégration bancaire. VeriPass prépare un dossier KYC et une décision de revue, mais l'ouverture comptable effective, l'activation commerciale, les échanges avec les systèmes de production et les processus de support devront être cadrés avec les directions concernées.",
                "Une autre limite concerne l'inclusion numérique. Un parcours mobile peut simplifier l'entrée en relation pour certains clients, mais il peut aussi créer des difficultés pour les personnes moins à l'aise avec la capture de documents ou la biométrie. L'industrialisation devra donc prévoir des chemins d'assistance, des messages simples et des reprises en agence lorsque le numérique ne suffit pas.",
                "La perspective la plus immédiate consiste à transformer les preuves techniques en indicateurs de pilotage. Il faudra mesurer le temps moyen de constitution, le taux de dossiers incomplets, le nombre de corrections OCR, les motifs de demande de complément, les cas de rejet biométrique et la charge par rôle back office. Ces indicateurs permettront de juger la solution sur son utilité réelle, et non sur la seule existence des composants techniques.",
                "Les perspectives portent donc sur l'industrialisation progressive : pilote BICEC, revue conformité, durcissement sécurité, mesure des indicateurs, extension des contrôles AML/CFT, amélioration de l'expérience mobile et possibilité de généraliser le socle à d'autres institutions de la zone CEMAC.",
            ],
        ),
        (
            "5.9 Synthèse du chapitre",
            [
                "Ce chapitre a évalué VeriPass comme prototype d'ingénierie. Les preuves disponibles montrent que le parcours peut être exécuté, que le pipeline OCR fonctionne sur un lot de tests, que la biométrie est intégrée comme signal de revue, et que l'architecture prépare la sécurité et l'auditabilité attendues dans un contexte KYC.",
                "La limite principale est claire : le prototype doit encore être confronté à un pilote réel avant toute affirmation de performance ou de rentabilité. Cette prudence renforce le mémoire, car elle distingue les résultats démontrés, les choix justifiés et les travaux nécessaires pour passer d'un prototype robuste à une solution bancaire industrialisée.",
            ],
        ),
    ]
    replace_block_between(
        doc,
        "Chapitre 5 : Évaluation, sécurité, déploiement, impact et limites",
        "Conclusion générale",
        ch5_sections,
    )

    conclusion = [
        "Ce mémoire avait pour objectif de concevoir et développer BICEC VeriPass, une plateforme d'onboarding KYC destinée à digitaliser la constitution du dossier client, structurer les contrôles documentaires et biométriques, puis préparer la revue humaine dans un cadre bancaire conforme. La problématique de départ portait sur un équilibre difficile : accélérer l'entrée en relation sans affaiblir l'identification, la conservation des preuves, la sécurité et l'auditabilité exigées par le métier bancaire.",
        "Le travail réalisé montre qu'un parcours KYC digital ne peut pas être réduit à un formulaire en ligne. Il doit organiser une chaîne complète : authentification du client, capture des pièces, extraction OCR, correction des champs, preuve de vie, comparaison faciale, consentement, signature, soumission, revue back office, décision motivée et journalisation. L'apport de VeriPass est d'avoir assemblé ces éléments dans une architecture cohérente, fondée sur une PWA mobile, une API FastAPI, PostgreSQL, Redis, Celery, un stockage documentaire et un back office de décision.",
        "Sur le plan technique, le projet a permis de mettre en pratique des compétences d'ingénierie logiciel et data. La conception a exigé de modéliser les entités KYC, de définir des états de dossier, de relier les documents à leurs métadonnées, de traiter les sorties OCR comme des données incertaines et de conserver les traces nécessaires à l'audit. La réalisation a également renforcé la maîtrise des contrats API, de l'orchestration asynchrone, des conteneurs Docker, des files Celery et de la séparation entre collecte, traitement et décision.",
        "Sur le plan méthodologique, le projet a imposé une discipline de traçabilité. Chaque besoin devait être relié à un acteur, chaque acteur à une action, chaque action à une donnée, et chaque donnée à une preuve ou à un contrôle. Cette logique a évité de réduire le mémoire à une application de démonstration. Elle a replacé le développement dans une logique d'ingénierie bancaire, où la valeur d'une fonctionnalité dépend de sa capacité à être comprise, contrôlée et justifiée.",
        "Sur le plan organisationnel, le stage a montré que la qualité d'une solution bancaire ne dépend pas seulement de sa performance technique. Elle dépend aussi de sa capacité à s'intégrer aux rôles existants, à respecter les responsabilités des agents, à rendre les anomalies compréhensibles et à produire des preuves exploitables. Les échanges avec l'encadrement BICEC ont donc orienté le projet vers une automatisation assistée, et non vers une décision automatique.",
        "Les résultats obtenus restent ceux d'un prototype avancé. Les preuves de parcours et de traitement OCR confirment la faisabilité technique, mais elles ne remplacent pas un pilote de production. Les seuils OCR et biométriques, les performances en conditions réelles, l'accessibilité du parcours mobile, la gouvernance des données biométriques et l'intégration aux systèmes bancaires devront faire l'objet d'une validation complémentaire.",
        "Les perspectives d'industrialisation sont néanmoins claires. Elles concernent le durcissement sécurité, la validation conformité, la supervision, la calibration sur données représentatives, l'amélioration de l'expérience mobile, l'extension des contrôles AML/CFT et la mesure des gains opérationnels. À terme, VeriPass peut devenir un socle RegTech pour l'onboarding bancaire dans la zone CEMAC, à condition de conserver son principe directeur : automatiser ce qui peut être structuré, journaliser ce qui doit être prouvé, et laisser aux acteurs habilités la responsabilité des décisions sensibles.",
    ]
    replace_after_heading_until(doc, "Conclusion générale", "Références bibliographiques", conclusion)

    # Bibliography cleanup with verified references and fewer invented entries.
    bibliography = [
        "Banque Internationale du Cameroun pour l'Epargne et le Crédit. (2026). Présentation institutionnelle de la BICEC. https://www.bicec.com",
        "Bulatov, K. B., Bezmaternykh, P. V., Nikolaev, D. P., & Arlazarov, V. V. (2022). Towards a unified framework for identity documents analysis and recognition. Computer Optics, 46(3), 436-454. https://doi.org/10.18287/2412-6179-CO-1024",
        "Carta, S., Giuliani, A., Piano, L., & Tiddia, S. G. (2024). An end-to-end OCR-free solution for identity document information extraction. Procedia Computer Science, 246, 453-462. https://doi.org/10.1016/j.procs.2024.09.425",
        "Commission Bancaire de l'Afrique Centrale. (2023). Règlement COBAC R-2023/01 relatif aux diligences des établissements assujettis en matière de lutte contre le blanchiment des capitaux, le financement du terrorisme et de la prolifération.",
        "Consultative Group to Assist the Poor. (2021). Regulation for inclusive digital finance. CGAP.",
        "Droit Médias Finance. (2024). Le nouveau règlement COBAC R-2023/01 et les diligences KYC des établissements assujettis.",
        "Du, Y., Li, C., Guo, R., Yin, X., Liu, W., Zhou, J., Bai, Y., Yu, Z., Yang, Y., Dang, Q., & Wang, H. (2020). PP-OCR: A practical ultra lightweight OCR system. arXiv. https://arxiv.org/abs/2009.09941",
        "Financial Action Task Force. (2020). Guidance on digital identity. FATF. https://www.fatf-gafi.org",
        "FastAPI. (2026). FastAPI documentation. https://fastapi.tiangolo.com/",
        "Docker. (2026). Docker Compose documentation. https://docs.docker.com/compose/",
        "OpenAPI Initiative. (2024). OpenAPI Specification. https://spec.openapis.org",
        "PostgreSQL Global Development Group. (2026). PostgreSQL documentation. https://www.postgresql.org/docs/",
        "Yu, Z., Qin, Y., Li, X., Zhao, C., Lei, Z., & Zhao, G. (2023). Deep learning for face anti-spoofing: A survey. IEEE Transactions on Pattern Analysis and Machine Intelligence, 45(5), 5609-5631.",
    ]
    replace_after_heading_until(doc, "Références bibliographiques", "Annexes", bibliography)
    for p in doc.paragraphs:
        if p.text.strip() in bibliography:
            p.style = "Bibliographie"

    # Rebuild preliminary lists after all renumbering and table removals.
    replace_between_paragraphs(
        doc,
        "Liste des tableaux",
        "Liste des figures",
        [
            "Tableau 1. Abréviations utilisées",
            "Tableau 2. Fiche signalétique synthétique de la BICEC et du stage",
            "Tableau 3. Limites du processus manuel et réponses du pipeline numérique",
            "Tableau 4. Architecture logique et responsabilités techniques",
            "Tableau 5. Gestion des échecs OCR et biométriques",
        ],
        "Front Matter",
    )
    replace_between_paragraphs(
        doc,
        "Liste des figures",
        "LISTE DES ABRÉVIATIONS",
        [
            "Figure 1. Architecture logique simplifiée de BICEC VeriPass",
            "Figure 2. Pipeline de données et états principaux du dossier KYC",
            "Figure 3. MLD simplifié des entités KYC principales",
        ],
        "Front Matter",
    )
    replace_paragraph(doc, "Tableau 2. Abréviations utilisées", "Tableau 1. Abréviations utilisées")
    replace_paragraph(doc, "Tableau 3. Fiche signalétique synthétique de la BICEC et du stage", "Tableau 2. Fiche signalétique synthétique de la BICEC et du stage")
    replace_paragraph(doc, "Tableau 4. Limites du processus manuel et réponses attendues du pipeline numérique", "Tableau 3. Limites du processus manuel et réponses attendues du pipeline numérique")
    replace_paragraph(doc, "Tableau 5. Architecture logique et responsabilités techniques", "Tableau 4. Architecture logique et responsabilités techniques")
    replace_paragraph(doc, "Tableau 6. Gestion des échecs OCR et biométriques", "Tableau 5. Gestion des échecs OCR et biométriques")
    replace_paragraph(doc, "Figure 2. Architecture logique simplifiée de BICEC VeriPass", "Figure 1. Architecture logique simplifiée de BICEC VeriPass")
    replace_paragraph(doc, "Figure 3. Pipeline de données et états principaux du dossier KYC", "Figure 2. Pipeline de données et états principaux du dossier KYC")
    replace_paragraph(doc, "Figure 4. MLD simplifié des entités KYC principales", "Figure 3. MLD simplifié des entités KYC principales")

    annex_replacements = {
        "La fiche de validation indique que le stage se déroule à la BICEC, à Douala, sous la responsabilité de M. KOMBE LELE Jackson Parfait, responsable du département Étude et Développement. Le sujet validé est : Conception et Développement d'un écosystème intelligent d'acquisition client : intégration de mécanismes avancés de vérification pour l'automatisation de la conformité KYC. L'objectif poursuivi est de digitaliser l'onboarding client pour réduire le délai d'enrôlement et automatiser le contrôle documentaire KYC. La période de stage indiquée s'étend du 19 janvier 2026 au 18 juin 2026.": "La fiche de validation indique que le stage se déroule à la BICEC, à Douala, sous la responsabilité de M. KOMBE LELE Jackson Parfait, responsable du département Étude et Développement. Le sujet validé est : Conception et Développement d'un écosystème intelligent d'acquisition client : intégration de mécanismes avancés de vérification pour l'automatisation de la conformité KYC. L'objectif poursuivi est de digitaliser l'onboarding client pour réduire le délai d'enrôlement et automatiser le contrôle documentaire KYC. La période de stage indiquée s'étend du 19 janvier 2026 au 18 juin 2026. La fiche mentionne un effectif d'environ 600 collaborateurs.",
        "Endpoint de référence : POST /api/v1/kyc/session/start crée ou reprend une session éditable. Le flux de soumission s'appuie ensuite sur POST /api/v1/kyc/capture/cni, POST /api/v1/kyc/ocr/review, POST /api/v1/kyc/liveness/submit, POST /api/v1/kyc/consent/submit et POST /api/v1/kyc/submit. La réponse de soumission contient session_id, status, message et access_level.": "Endpoint de référence : POST /api/v1/kyc/session/start crée ou reprend une session éditable. Le flux de soumission s'appuie ensuite sur POST /api/v1/kyc/capture/cni, POST /api/v1/kyc/ocr/review, POST /api/v1/kyc/liveness/submit, POST /api/v1/kyc/address/submit, POST /api/v1/kyc/consent/submit, POST /api/v1/kyc/signature/submit, GET /api/v1/kyc/readiness et POST /api/v1/kyc/submit. La réponse de soumission contient notamment l'identifiant de session, le statut, le message et le niveau d'accès.",
        "Les captures d'écran disponibles dans docs/test-evidence/latest illustrent les files back office, les alertes AML, la démonstration de conformité, les écrans mobiles du parcours KYC et les preuves d'exécution E2E. Elles seront sélectionnées et insérées dans la version finale des chapitres 4 et 5.": "Les captures d'écran disponibles dans docs/test-evidence/latest illustrent les files back office, les alertes AML, la démonstration de conformité, les écrans mobiles du parcours KYC et les preuves d'exécution de bout en bout. Elles constituent des preuves complémentaires pour la soutenance et peuvent être sélectionnées dans les annexes détaillées si le volume final du mémoire le permet.",
        "Tables principales : kyc_sessions pour le dossier et ses états, documents pour les pièces et leurs hash, ocr_fields pour les champs extraits et corrigés, biometric_results pour les scores liveness et face matching, validation_decisions pour les décisions back office, audit_logs pour la traçabilité transversale.": "Tables principales : kyc_sessions pour le dossier et ses états, documents pour les pièces et leurs hash, ocr_fields pour les champs extraits et corrigés, biometric_results pour les scores liveness et face matching, validation_decisions pour les décisions back office, aml_alerts pour les alertes de conformité et audit_log pour la traçabilité transversale.",
    }
    for old, new in annex_replacements.items():
        replace_paragraph(doc, old, new)

    replace_paragraph(
        doc,
        "Services principaux : nginx comme proxy TLS, pwa pour l'application mobile, back office pour l'interface interne, api pour FastAPI, postgres pour la base, redis pour le broker, celery_ocr pour les traitements OCR, celery_notifications pour les notifications, celery_beat pour les tâches planifiées, flower pour la supervision Celery. Les volumes externes principaux sont documents_storage, db_storage et db_backups.",
        "Services principaux : nginx comme proxy TLS, pwa pour l'application mobile, back office pour l'interface interne, api pour FastAPI, postgres pour la base, redis pour le broker, celery_ocr pour les traitements OCR, celery_notifications pour les notifications, celery_beat pour les tâches planifiées, flower pour la supervision Celery. Les volumes externes principaux sont documents_storage, db_storage, db_backups et models_storage. Cette composition sépare l'interface, les règles métier, la persistance relationnelle, les fichiers KYC, les traitements différés et la supervision. Elle constitue une base de démonstration reproductible, mais une production demanderait une politique de secrets, de sauvegardes, de supervision et de durcissement plus complète.",
    )
    replace_paragraph(
        doc,
        "Endpoint de référence : POST /api/v1/kyc/session/start crée ou reprend une session éditable. Le flux de soumission s'appuie ensuite sur POST /api/v1/kyc/capture/cni, POST /api/v1/kyc/ocr/review, POST /api/v1/kyc/liveness/submit, POST /api/v1/kyc/address/submit, POST /api/v1/kyc/consent/submit, POST /api/v1/kyc/signature/submit, GET /api/v1/kyc/readiness et POST /api/v1/kyc/submit. La réponse de soumission contient notamment l'identifiant de session, le statut, le message et le niveau d'accès.",
        "Endpoint de référence : POST /api/v1/kyc/session/start crée ou reprend une session éditable. Le flux de soumission s'appuie ensuite sur POST /api/v1/kyc/capture/cni, POST /api/v1/kyc/ocr/review, POST /api/v1/kyc/liveness/submit, POST /api/v1/kyc/address/submit, POST /api/v1/kyc/consent/submit, POST /api/v1/kyc/signature/submit, GET /api/v1/kyc/readiness et POST /api/v1/kyc/submit. La réponse de soumission contient notamment l'identifiant de session, le statut, le message et le niveau d'accès. Ce contrat illustre la philosophie du projet : chaque étape porte une responsabilité claire, produit une preuve et peut être reprise ou auditée sans confondre la collecte client avec la décision bancaire.",
    )
    replace_paragraph(
        doc,
        "Tables principales : kyc_sessions pour le dossier et ses états, documents pour les pièces et leurs hash, ocr_fields pour les champs extraits et corrigés, biometric_results pour les scores liveness et face matching, validation_decisions pour les décisions back office, aml_alerts pour les alertes de conformité et audit_log pour la traçabilité transversale.",
        "Tables principales : kyc_sessions pour le dossier et ses états, documents pour les pièces et leurs hash, ocr_fields pour les champs extraits et corrigés, biometric_results pour les scores liveness et face matching, validation_decisions pour les décisions back office, aml_alerts pour les alertes de conformité et audit_log pour la traçabilité transversale. La séparation de ces tables matérialise la séparation des responsabilités : un document n'est pas un champ OCR, un score n'est pas une décision, et une décision n'est pas un journal d'audit. Cette distinction est essentielle pour expliquer le dossier lors d'un contrôle.",
    )
    for p in list(doc.paragraphs):
        t = p.text.strip()
        if t.endswith("une politique de secrets, de sauvegardes, de supervision et de durcissement plus complète."):
            cur = insert_after(
                p,
                "Dans une lecture d'exploitation, l'annexe B montre aussi les frontières du prototype. Le service API concentre les règles métier, PostgreSQL porte la cohérence transactionnelle, Redis porte les files et les verrous, tandis que les workers exécutent les tâches qui ne doivent pas bloquer l'utilisateur. Cette séparation facilite le diagnostic : une panne de worker OCR n'a pas le même effet qu'une indisponibilité de l'API ou qu'une saturation de la base.",
                "Memoire Body",
            )
            insert_after(
                cur,
                "Pour une mise en production, cette composition devrait être complétée par une gestion stricte des secrets, une rotation des certificats, des sauvegardes testées, une supervision centralisée, une politique de rétention et une procédure de reprise. L'annexe ne prétend donc pas être un manuel d'exploitation complet; elle donne le périmètre technique à partir duquel ces procédures peuvent être construites.",
                "Memoire Body",
            )
        if t.endswith("sans confondre la collecte client avec la décision bancaire."):
            cur = insert_after(
                p,
                "Le contrat API permet également de comprendre la responsabilité de chaque réponse. Une réponse de capture confirme la réception d'une pièce et son rattachement au dossier; une réponse de readiness indique si le dossier peut être transmis; une réponse de soumission modifie le statut de cycle de vie. Cette distinction empêche de traiter toutes les réponses comme de simples confirmations techniques.",
                "Memoire Body",
            )
            insert_after(
                cur,
                "Dans un projet bancaire, cette précision a une conséquence directe sur l'audit. Si un client conteste une décision ou si un contrôle interne demande la reconstitution du dossier, le système doit pouvoir montrer quelles étapes ont été effectuées, dans quel ordre, avec quelles données et par quel acteur. Les endpoints décrits dans l'annexe C structurent cette chronologie.",
                "Memoire Body",
            )
        if t.endswith("Cette distinction est essentielle pour expliquer le dossier lors d'un contrôle."):
            cur = insert_after(
                p,
                "Le schéma prépare aussi l'analyse des erreurs. Si un champ OCR est faux, la correction peut être étudiée sans remettre en cause tout le document. Si un score biométrique est faible, il peut être comparé à la décision humaine sans modifier la pièce d'origine. Si une décision est contestée, le journal d'audit peut reconstituer l'acteur, l'action et le moment. Cette granularité est l'une des différences majeures entre un simple stockage de fichiers et un vrai système KYC.",
                "Memoire Body",
            )
            insert_after(
                cur,
                "La base relationnelle joue donc un rôle de mémoire opérationnelle. Elle conserve les liens entre dossier, pièces, champs, scores, alertes et décisions. Cette mémoire est indispensable pour améliorer le système après les premiers usages, car elle permet d'identifier les erreurs récurrentes, les champs les plus souvent corrigés, les motifs de demande de complément et les points du parcours qui demandent une assistance client.",
                "Memoire Body",
            )

    # Page breaks for major parts.
    style_major_headings(doc)
    for p in doc.paragraphs:
        if p.text.strip() == "TABLE DES MATIÈRES":
            insert_toc_after(p)
            break

    # Basic text hygiene.
    for p in doc.paragraphs:
        if p.text:
            text = p.text
            text = text.replace("Know Your Customer", "Know Your Customer")
            text = text.replace("anti spoofing", "anti spoofing")
            text = text.replace("back-office", "back office")
            if text != p.text:
                set_text(p, text)

    doc.save(OUT)

    # Guardrails requested by the author.
    d2 = Document(OUT)
    all_text = "\n".join(p.text for p in d2.paragraphs)
    for table in d2.tables:
        for row in table.rows:
            for cell in row.cells:
                all_text += "\n" + cell.text
    bad = {
        "em_dash": "—" in all_text,
        "en_dash": "–" in all_text,
        "spaced_hyphen": " - " in all_text,
        "Jean": "Jean" in all_text,
        "Thomas": "Thomas" in all_text,
        "Sylvie": "Sylvie" in all_text,
        "Admin IT": "Admin IT" in all_text,
        "MiniFASNet": "MiniFASNet" in all_text,
        "5 000": "5 000" in all_text,
        "ROI": "ROI" in all_text,
        "15 minutes": "15 minutes" in all_text,
    }
    print(f"Wrote {OUT}")
    print("Guardrails:", bad)


if __name__ == "__main__":
    main()
